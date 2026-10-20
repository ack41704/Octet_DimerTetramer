"""
Rate constants for the 9 EphrinB3-Fc extended runs (15 h dissociation): Global Fits 69-75, 77 (PBST) and 65 (2026.96, PBSTB).
Numbering as in the master table: 1 = tetramer, 2 = dimer.

  Ka2 (1/Ms)    slope of k_obs vs [EphB2] over all association curves (pipeline k_obs regression)
  Kdis2, Kdis1  fast / slow rate of the global biexponential dissociation fit (all fitted sensors)
  Ka1 (1/Ms)    top concentration only: its association trace refit alone (1500 evenly spaced points,
                fast + slow exponential); the slow phase is treated as a second, bimolecular binding step,
                k_slow = Ka1*[EphB2] + Kdis1  ->  Ka1 = (k_slow - Kdis1) / [EphB2]_top.
                Assumes the slow rate is linear in [EphB2]; only one concentration shows the phase, so this
                is assumed, not tested.
  KD1, KD2      Kdis/Ka, in nM
  R^2           association, top concentration (Ka1 fit); association, all curves pooled (Ka2);
                k_obs regression (Ka2); dissociation, global fit (Kdis1, Kdis2)

Outputs (analysis folder): EB3_extended_constants_9-30-26.xlsx, EB3_extended_sensorgrams_9-30-26.pdf,
EB3_extended_sensorgrams_overview.png. Workbook formulas are verified here by an independent evaluator.
"""
from paths import DATA_DIR, WORK_DIR  # noqa: E402,F401  (repo copy: portable data / work locations)
import os
import re
import sys
import math
import pickle

import numpy as np
from scipy import stats
from scipy.optimize import curve_fit
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
from mpl_toolkits.axes_grid1.inset_locator import inset_axes

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from batch_fit import sub, f2a  # noqa: E402

W = WORK_DIR
OUT_XLSX = os.path.join(W, "EB3_extended_constants_9-30-26.xlsx")
OUT_PDF = os.path.join(W, "EB3_extended_sensorgrams_9-30-26.pdf")
OUT_PNG = os.path.join(W, "EB3_extended_sensorgrams_overview.png")
RUNS = [("2026.101", 69, "EphB2-WT"), ("2026.103", 70, "EphB2-WT"), ("2026.108", 73, "EphB2-WT"),
        ("2026.106", 71, "EphB2-R79P"), ("2026.107", 72, "EphB2-R79P"), ("2026.109", 74, "EphB2-R79P"),
        ("2026.111", 75, "EphB2-His"), ("2026.116", 77, "EphB1-His"), ("2026.96 (PBSTB)", 65, "EphB2-His")]
# Ka1 needs a sharp first (dimer) phase plus a separate slow phase at the top concentration. If the "slow" component
# carries most of the signal, it is the dimer binding itself (2026.96 in PBSTB: slow k = that concentration's k_obs).
KA1_MAX_SLOW_SHARE = 0.5
INK, INK2, GRID, GREY = "#1f1f1f", "#5a5a58", "#e6e6e3", "#9A9A96"


def two(t, A1, k1, A2, k2):
    return A1 * (1 - np.exp(-k1 * t)) + A2 * (1 - np.exp(-k2 * t))


def load():
    cache = pickle.load(open(os.path.join(W, "fit_cache.pkl"), "rb"))
    by = {e["r"]["experiment"]: e for e in cache.values()
          if "error" not in e and e["r"]["interaction"].startswith("EphrinB3-Fc")}
    runs = []
    for exp, gf, analyte in RUNS:
        e = by[exp]; r, s = e["r"], e["s"]; tb = r["tb"]
        # ---- Ka1: top-concentration trace refit alone
        top = max(w["conc_nM"] for w in r["keep"])
        w = max([w for w in r["keep"] if w["conc_nM"] == top], key=lambda w: w["y"].max())
        m = w["t"] <= tb
        t, y = w["t"][m], w["y"][m]
        i = np.unique(np.linspace(0, len(t) - 1, 1500).astype(int)); t, y = t[i], y[i]
        p, _ = curve_fit(two, t, y, p0=[y[-1] * .75, 5e-3, y[-1] * .25, 4e-4],
                         bounds=([0, 1e-4, 0, 1e-6], [5, 1, 5, 5e-3]), maxfev=40000)
        A1, k1, A2, k2 = p
        if k1 < k2:
            A1, k1, A2, k2 = A2, k2, A1, k1
        r2_top = 1 - np.sum((y - two(t, A1, k1, A2, k2)) ** 2) / np.sum((y - y.mean()) ** 2)
        # ---- association R^2 over all fitted curves (the per-sensor fits behind the k_obs regression)
        ss_res = ss_tot = 0.0; ally = []
        per = []
        for ww in r["keep"]:
            if ww["conc_nM"] not in r["assoc_fits"]:
                continue
            ma = ww["t"] <= tb
            ya, yh = ww["y"][ma], f2a(ww["t"][ma], *r["assoc_fits"][ww["conc_nM"]])
            ally.append(ya); ss_res += np.sum((ya - yh) ** 2)
            per.append(1 - np.sum((ya - yh) ** 2) / np.sum((ya - ya.mean()) ** 2))
        ally = np.concatenate(ally)
        r2_all = 1 - ss_res / np.sum((ally - ally.mean()) ** 2)
        runs.append(dict(exp=exp, gf=gf, analyte=analyte, e=e, top=top, top_well=w, fast=(A1, k1), slow=(A2, k2),
                         k_slow=k2, kdis1=r["ks"], kdis2=r["kf"], ka2=r["ka1"],
                         ka1=(k2 - r["ks"]) / (top * 1e-9), r2_top=float(r2_top), r2_all=float(r2_all),
                         r2_assoc_median=float(np.median(per)), r2_kobs=r["r2k"], r2_dis=r["R2d"],
                         assoc_min=s["assoc_min"], dis_h=s["dissoc_h"], n_fit=len(r["keep"]), buffer=s["buffer"],
                         n_all=len(e["mask"]), slow_share=A2 * (1 - np.exp(-k2 * tb)) / two(tb, A1, k1, A2, k2)))
    for x in runs:
        x["ka1_ok"] = x["slow_share"] <= KA1_MAX_SLOW_SHARE
        x["KD1_nM"] = x["kdis1"] / x["ka1"] * 1e9
        x["KD2_nM"] = x["kdis2"] / x["ka2"] * 1e9
    return runs


# ------------------------------------------------------------------ figures
def sensorgram(ax, x, inset=True, legend=True, small=False):
    e = x["e"]; r = e["r"]; tb = r["tb"]
    concs = sorted({w["conc_nM"] for w in r["keep"]})
    cols = dict(zip(concs, plt.cm.viridis(np.linspace(0.05, 0.9, len(concs)))))
    end = r["keep"][0]["t"][-1]

    def draw(a, lw=1.3, ms=1.6):
        for w, m in zip(e["parsed"]["wells"], e["mask"]):
            if m:
                continue
            ta, ya = sub(w["t"], w["y"], 200, 200)
            a.plot(ta / 3600, ya, ".", ms=ms * 0.7, color=GREY, alpha=0.4)
        per_well = x.get("assoc_by_well")          # per-sensor fits (needed when two sensors share a concentration)
        for j, (w, A, B) in enumerate(zip(r["keep"], r["A1s"], r["A2s"])):
            c = cols[w["conc_nM"]]
            ma, md = w["t"] <= tb, w["t"] >= tb
            ta, ya = sub(w["t"][ma], w["y"][ma], 150, 120)
            td, yd = sub(w["t"][md] - tb, w["y"][md], 200, 150)
            a.plot(np.r_[ta, td + tb] / 3600, np.r_[ya, yd], ".", ms=ms, alpha=0.35, color=c)
            pa = per_well[j] if per_well is not None else r["assoc_fits"].get(w["conc_nM"])
            if pa is not None:
                tt = np.linspace(0, tb, 300)
                a.plot(tt / 3600, f2a(tt, *pa), "-", lw=lw, color=c)
            tt = np.r_[0, np.geomspace(1, end - tb, 400)]
            a.plot((tt + tb) / 3600, A * np.exp(-r["kf"] * tt) + B * np.exp(-r["ks"] * tt), "-", lw=lw, color=c,
                   label=f"{w['conc_nM']:g} nM")
        a.axvline(tb / 3600, color=INK2, ls=":", lw=0.9)
    draw(ax)
    ys = np.concatenate([w["y"] for w in r["keep"]]); hi = float(ys.max())
    ax.set_ylim(min(0, float(np.percentile(ys, 0.1))) - 0.05 * hi, hi * 1.12)
    ax.set_xlim(-0.2, end / 3600 + 0.2)
    ax.text(tb / 3600, hi * 1.08, "  association | dissociation", fontsize=6.5 if small else 7.5, color=INK2, va="center")
    ax.set_xlabel("time (h)", fontsize=8 if small else 9); ax.set_ylabel("response (nm)", fontsize=8 if small else 9)
    ax.yaxis.grid(True, color=GRID, lw=0.8); ax.set_axisbelow(True)
    for sp in ("top", "right"):
        ax.spines[sp].set_visible(False)
    ax.tick_params(labelsize=7 if small else 8, colors=INK2)
    if legend:
        h, l = ax.get_legend_handles_labels()
        ax.legend(h, l, fontsize=6.5, loc="upper right", frameon=False, bbox_to_anchor=(1.0, 0.93), ncol=1)
    if inset:
        ins = inset_axes(ax, width="36%", height="42%", loc="center right", bbox_to_anchor=(0, -0.08, 0.98, 1),
                         bbox_transform=ax.transAxes, borderpad=0.3)
        draw(ins, lw=1.0, ms=1.2)
        ins.set_xlim(-0.05, tb / 3600 + 1.0); ins.set_ylim(ax.get_ylim())
        ins.set_title("first 2.5 h", fontsize=7, color=INK2, pad=2)
        ins.tick_params(labelsize=6, colors=INK2)
        for sp in ("top", "right"):
            ins.spines[sp].set_visible(False)


def run_page(pdf, x, k, n):
    fig = plt.figure(figsize=(11, 8.5))
    fig.text(0.05, 0.955, f"EphrinB3-Fc + {x['analyte']}   —   {x['exp']}   (Global Fit # {x['gf']})",
             fontsize=13, weight="bold", color=INK)
    fig.text(0.05, 0.928, f"{x['buffer']} · association {x['assoc_min']:.0f} min · dissociation {x['dis_h']:.1f} h · "
             f"{x['n_fit']} of {x['n_all']} sensors fitted · dots = data, lines = fits (grey = excluded sensors)",
             fontsize=9, color=INK2)
    fig.text(0.95, 0.955, f"page {k} / {n}", fontsize=8, color="#777", ha="right")
    ax = fig.add_axes([0.07, 0.38, 0.6, 0.52])
    sensorgram(ax, x, inset=False)
    axz = fig.add_axes([0.73, 0.38, 0.24, 0.52])
    sensorgram(axz, x, inset=False, legend=False, small=True)
    axz.set_xlim(-0.05, x["e"]["r"]["tb"] / 3600 + 1.0)
    axz.set_title("first 2.5 h (zoom)", fontsize=9, color=INK2, loc="left")
    ka1_txt, kd1_txt = (f"{x['ka1']:.3e}", f"{x['KD1_nM']:.3g}") if x["ka1_ok"] else ("n.d.", "n.d.")
    ka1_note = (f"{x['top']:g} nM slow phase: k = {x['k_slow']:.2e} 1/s ({100 * x['slow_share']:.0f}% of signal)"
                if x["ka1_ok"] else f"no separate sharp phase at {x['top']:g} nM: slow component ({100 * x['slow_share']:.0f}% "
                f"of signal) is the dimer binding")
    rows = [("Ka1  tetramer (1/Ms)", ka1_txt, ka1_note),
            ("Kdis1 tetramer (1/s)", f"{x['kdis1']:.3e}", f"half-life {math.log(2) / x['kdis1'] / 86400:.2f} days"),
            ("KD1  tetramer (nM)", kd1_txt, "Kdis1 / Ka1"),
            ("Ka2  dimer (1/Ms)", f"{x['ka2']:.3e}", "slope of k_obs vs [EphB2], all curves"),
            ("Kdis2 dimer (1/s)", f"{x['kdis2']:.3e}", f"half-life {math.log(2) / x['kdis2'] / 60:.1f} min"),
            ("KD2  dimer (nM)", f"{x['KD2_nM']:.3g}", "Kdis2 / Ka2"),
            ("R² association, top conc", f"{x['r2_top']:.4f}", "fit used for Ka1"),
            ("R² association, all curves", f"{x['r2_all']:.4f}", f"used for Ka2 (k_obs regression R² {x['r2_kobs']:.3f})"),
            ("R² dissociation", f"{x['r2_dis']:.4f}", "global fit, all curves (Kdis1, Kdis2)")]
    y = 0.29
    for name, val, note in rows:
        fig.text(0.07, y, name, fontsize=9, color=INK, family="DejaVu Sans Mono")
        fig.text(0.36, y, val, fontsize=9, color=INK, family="DejaVu Sans Mono", weight="bold")
        fig.text(0.5, y, note, fontsize=8, color=INK2)
        y -= 0.028
    pdf.savefig(fig); plt.close(fig)


def table_page(pdf, runs):
    fig = plt.figure(figsize=(11, 8.5))
    fig.text(0.04, 0.95, "EphrinB3-Fc extended runs: rate constants and fit quality", fontsize=13, weight="bold", color=INK)
    fig.text(0.04, 0.925, "1 = tetramer, 2 = dimer. Ka1 from the top-concentration slow phase, "
             "Ka1 = (k_slow − Kdis1) / [EphB2]top; Ka2 from all association curves; Kdis from the global dissociation fit.",
             fontsize=8.3, color=INK2)
    ax = fig.add_axes([0.02, 0.44, 0.96, 0.46]); ax.axis("off")
    head = ["GF #", "Run", "Analyte", "Ka1\n(1/Ms)", "Kdis1\n(1/s)", "KD1\n(nM)", "Ka2\n(1/Ms)", "Kdis2\n(1/s)",
            "KD2\n(nM)", "R² assoc\ntop conc", "R² assoc\nall curves", "R²\ndissoc"]
    rows = [[str(x["gf"]), x["exp"].split()[0], x["analyte"] + ("" if x["buffer"] == "PBST" else f" ({x['buffer']})"),
             f"{x['ka1']:.3e}" if x["ka1_ok"] else "n.d.", f"{x['kdis1']:.3e}",
             f"{x['KD1_nM']:.3g}" if x["ka1_ok"] else "n.d.",
             f"{x['ka2']:.3e}", f"{x['kdis2']:.3e}", f"{x['KD2_nM']:.3g}", f"{x['r2_top']:.4f}", f"{x['r2_all']:.4f}",
             f"{x['r2_dis']:.4f}"] for x in runs]
    wd = [0.05, 0.075, 0.1, 0.09, 0.09, 0.06, 0.09, 0.09, 0.06, 0.08, 0.08, 0.07]
    tb = ax.table(cellText=rows, colLabels=head, loc="upper center", cellLoc="center", colWidths=[w / sum(wd) for w in wd])
    tb.auto_set_font_size(False); tb.set_fontsize(7.8); tb.scale(1, 1.8)
    for (i, j), c in tb.get_celld().items():
        c.set_edgecolor(GRID)
        if i == 0:
            c.set_facecolor("#f2f2ef"); c.set_text_props(weight="bold", color=INK)
    fig.text(0.04, 0.38, "EphB2-WT vs EphB2-R79P (n = 3 each), Student's t-test, two-tailed, equal variance",
             fontsize=10.5, weight="bold", color=INK)
    ax2 = fig.add_axes([0.02, 0.07, 0.96, 0.29]); ax2.axis("off")
    rows2 = []
    for key, name, f in (("ka1", "Ka1 (1/Ms)", "{:.3e}"), ("kdis1", "Kdis1 (1/s)", "{:.3e}"), ("KD1_nM", "KD1 (nM)", "{:.3g}"),
                         ("ka2", "Ka2 (1/Ms)", "{:.3e}"), ("kdis2", "Kdis2 (1/s)", "{:.3e}"), ("KD2_nM", "KD2 (nM)", "{:.3g}")):
        wt = [x[key] for x in runs if x["analyte"] == "EphB2-WT"]; mu = [x[key] for x in runs if x["analyte"] == "EphB2-R79P"]
        p = stats.ttest_ind(wt, mu, equal_var=True).pvalue
        rows2.append([name, f"{f.format(np.mean(wt))} ± {f.format(np.std(wt, ddof=1))}",
                      f"{f.format(np.mean(mu))} ± {f.format(np.std(mu, ddof=1))}", f"{np.mean(mu) / np.mean(wt):.2f}",
                      ("<0.001" if p < 0.001 else f"{p:.3f}")])
    tb2 = ax2.table(cellText=rows2, colLabels=["Constant", "WT mean ± SD", "R79P mean ± SD", "R79P / WT", "p"],
                    loc="upper center", cellLoc="center")
    tb2.auto_set_font_size(False); tb2.set_fontsize(8.2); tb2.scale(1, 1.75)
    for (i, j), c in tb2.get_celld().items():
        c.set_edgecolor(GRID)
        if i == 0:
            c.set_facecolor("#f2f2ef"); c.set_text_props(weight="bold", color=INK)
    fig.text(0.04, 0.03, "Ka1 assumes the slow phase is a bimolecular step (rate linear in [EphB2]); only the top "
             "concentration shows it, so that linearity is assumed, not tested.\n2026.96 is in PBSTB (not comparable with the "
             "PBST runs, not in the t-tests); its top-concentration trace has no separate sharp phase, so Ka1/KD1 are n.d.",
             fontsize=7.8, color=INK2, linespacing=1.5)
    pdf.savefig(fig); plt.close(fig)


def overview_png(runs):
    ncol = 3 if len(runs) == 9 else 4
    nrow = int(np.ceil(len(runs) / ncol))
    fig, axes = plt.subplots(nrow, ncol, figsize=(5.4 * ncol, 3.9 * nrow + 0.4))
    for ax in axes.ravel()[len(runs):]:
        ax.axis("off")
    for ax, x in zip(axes.ravel(), runs):
        sensorgram(ax, x, inset=False, legend=False, small=True)
        ax.set_title(f"{x['analyte']}  {x['exp']}  (GF {x['gf']})" + ("" if x["buffer"] == "PBST" or x["buffer"] in x["exp"]
                     else f"  {x['buffer']}"), fontsize=9.5, weight="bold", color=INK, loc="left")
    fig.suptitle("EphrinB3-Fc extended runs: full sensorgrams (association + 15 h dissociation), data = dots, fits = lines; "
                 "all PBST except 2026.96 (PBSTB)",
                 fontsize=12, weight="bold", color=INK, x=0.01, ha="left")
    fig.tight_layout(rect=[0, 0, 1, 0.95])
    fig.savefig(OUT_PNG, dpi=200); plt.close(fig)


# ------------------------------------------------------------------ workbook
def write_xlsx(runs):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    wb = Workbook()
    F, FB = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
    FILL = PatternFill("solid", fgColor="F2F2EF"); BOX = Border(bottom=Side(style="thin", color="D9D9D6"))

    def header(ws, cols, widths):
        for j, (h, w) in enumerate(zip(cols, widths), 1):
            c = ws.cell(row=1, column=j, value=h); c.font = FB; c.fill = FILL; c.border = BOX
            c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
            ws.column_dimensions[c.column_letter].width = w
        ws.row_dimensions[1].height = 46; ws.freeze_panes = "A2"

    ws = wb.active; ws.title = "Constants"
    cols = ["Global Fit #", "Run", "Analyte", "Top concentration (nM)", "Slow-phase rate at top conc, k_slow (1/s)",
            "Ka1 tetramer (1/Ms)", "Kdis1 tetramer (1/s)", "KD1 tetramer (nM)", "Ka2 dimer (1/Ms)", "Kdis2 dimer (1/s)",
            "KD2 dimer (nM)", "R² association, top concentration (Ka1)", "R² association, all curves (Ka2)",
            "R² k_obs vs [EphB2] regression (Ka2)", "R² dissociation, global fit (Kdis1, Kdis2)",
            "Tetramer half-life (days)", "Dimer half-life (min)", "Buffer", "Note"]
    header(ws, cols, [10, 15, 12, 12, 16, 14, 14, 11, 14, 14, 11, 15, 15, 15, 16, 12, 12, 9, 60])
    for i, x in enumerate(runs):
        rr = i + 2; x["row"] = rr
        vals = {1: x["gf"], 2: x["exp"], 3: x["analyte"], 4: x["top"], 5: float(f"{x['k_slow']:.10g}"),
                6: f"=(E{rr}-G{rr})/(D{rr}*1E-9)", 7: float(f"{x['kdis1']:.10g}"), 8: f"=G{rr}/F{rr}*1E9",
                9: float(f"{x['ka2']:.10g}"), 10: float(f"{x['kdis2']:.10g}"), 11: f"=J{rr}/I{rr}*1E9",
                12: float(f"{x['r2_top']:.10g}"), 13: float(f"{x['r2_all']:.10g}"), 14: float(f"{x['r2_kobs']:.10g}"),
                15: float(f"{x['r2_dis']:.10g}"), 16: f"=LN(2)/G{rr}/86400", 17: f"=LN(2)/J{rr}/60",
                18: x["buffer"], 19: ""}
        if not x["ka1_ok"]:
            vals[6] = vals[8] = "n.d."
            vals[19] = ("No separate sharp phase at the top concentration: the slow component "
                        f"({100 * x['slow_share']:.0f}% of the signal) has the same rate as that concentration's k_obs, "
                        "i.e. it is the dimer binding, so Ka1 and KD1 are not determinable. PBSTB: not comparable with PBST runs.")
        for j, v in vals.items():
            ws.cell(row=rr, column=j, value=v).font = F
        for col, fmt in ((4, "0.0"), (5, "0.000E+00"), (6, "0.000E+00"), (7, "0.000E+00"), (8, "0.000"), (9, "0.000E+00"),
                         (10, "0.000E+00"), (11, "0.00"), (12, "0.0000"), (13, "0.0000"), (14, "0.000"), (15, "0.0000"),
                         (16, "0.00"), (17, "0.0")):
            ws.cell(row=rr, column=col).number_format = fmt

    st = wb.create_sheet("WT vs R79P")
    header(st, ["Constant", "WT mean", "WT SD", "WT n", "R79P mean", "R79P SD", "R79P n", "R79P / WT",
                "Student's t-test p (two-tailed, equal variance)"], [22, 13, 13, 7, 13, 13, 8, 10, 18])
    wr = [x["row"] for x in runs if x["analyte"] == "EphB2-WT"]; mr = [x["row"] for x in runs if x["analyte"] == "EphB2-R79P"]
    assert wr == list(range(wr[0], wr[-1] + 1)) and mr == list(range(mr[0], mr[-1] + 1))
    for i, (name, c, fmt) in enumerate((("Ka1 (1/Ms)", "F", "0.000E+00"), ("Kdis1 (1/s)", "G", "0.000E+00"),
                                        ("KD1 (nM)", "H", "0.000"), ("Ka2 (1/Ms)", "I", "0.000E+00"),
                                        ("Kdis2 (1/s)", "J", "0.000E+00"), ("KD2 (nM)", "K", "0.00"))):
        rr = i + 2
        a, b = f"Constants!{c}{wr[0]}:{c}{wr[-1]}", f"Constants!{c}{mr[0]}:{c}{mr[-1]}"
        for j, v in enumerate([name, f"=AVERAGE({a})", f"=STDEV({a})", f"=COUNT({a})", f"=AVERAGE({b})",
                               f"=STDEV({b})", f"=COUNT({b})", f"=E{rr}/B{rr}", f"=TTEST({a},{b},2,2)"], 1):
            st.cell(row=rr, column=j, value=v).font = F
        for col in (2, 3, 5, 6):
            st.cell(row=rr, column=col).number_format = fmt
        st.cell(row=rr, column=8).number_format = "0.00"; st.cell(row=rr, column=9).number_format = "0.000"
    st.cell(row=9, column=1, value="WT = GF 69, 70, 73; R79P = GF 71, 72, 74. TTEST(range1, range2, 2, 2) = two-tailed, "
                                   "two-sample, equal-variance Student's t-test.").font = Font(name="Arial", size=9, italic=True)

    nt = wb.create_sheet("Notes"); nt.column_dimensions["A"].width = 125
    for i, t in enumerate([
        "EphrinB3-Fc extended runs (15 h dissociation): Global Fits 69-75 and 77 (PBST) and 65 (2026.96, PBSTB). Generated 2026-09-30. "
        "1 = tetramer, 2 = dimer. WT vs R79P uses the PBST runs only.",
        "Ka2: slope of k_obs vs [EphB2] from the per-sensor double-exponential association fits (all fitted curves).",
        "Kdis2, Kdis1: fast and slow rates of the global biexponential dissociation fit (shared rates, per-sensor amplitudes).",
        "Ka1: top-concentration association trace refit alone (1500 evenly spaced points), fast + slow exponential; the slow phase is "
        "treated as a bimolecular step, k_slow = Ka1*[EphB2] + Kdis1, so Ka1 = (k_slow - Kdis1)/[EphB2]_top (column F formula).",
        "Assumption: the slow rate is linear in [EphB2]. Only the top concentration shows a distinct slow phase, so this is assumed, not tested.",
        "KD = Kdis/Ka x 1e9 (nM). Half-lives = LN(2)/Kdis.",
        "R2 association (top conc) = the Ka1 fit; R2 association (all curves) = pooled over all fitted sensors' association fits; "
        "R2 dissociation = pooled global fit.",
        "Typed values are fit outputs (Constants D, E, G, I, J, L-O); all other numbers are formulas.",
        "Code: dimer_tetramer/eb3_extended_constants.py"], 1):
        c = nt.cell(row=i, column=1, value=t); c.font = F; c.alignment = Alignment(wrap_text=True)
    wb.save(OUT_XLSX)


def evaluate_workbook(path):
    """Independent evaluator for the formula subset used above (no LibreOffice here)."""
    from openpyxl import load_workbook
    from openpyxl.utils import column_index_from_string as ci, get_column_letter as gl
    wb = load_workbook(path); memo = {}

    def cell(sheet, ref):
        if (sheet, ref) not in memo:
            v = wb[sheet][ref].value
            memo[(sheet, ref)] = ev(sheet, v[1:]) if isinstance(v, str) and v.startswith("=") else v
        return memo[(sheet, ref)]

    def rng(sheet, c1, r1, c2, r2):
        return [cell(sheet, f"{gl(c)}{r}") for c in range(ci(c1), ci(c2) + 1) for r in range(int(r1), int(r2) + 1)]

    num = lambda v: [x for x in v if isinstance(x, (int, float)) and not isinstance(x, bool)]
    ns = dict(AVERAGE=lambda v: float(np.mean(num(v))), STDEV=lambda v: float(np.std(num(v), ddof=1)),
              COUNT=lambda v: len(num(v)), LN=math.log,
              TTEST=lambda a, b, tails, typ: float(stats.ttest_ind(num(a), num(b), equal_var=True).pvalue))
    pat = (r"(?:'([^']+)'!|([A-Za-z]+)!)?(?<![A-Za-z0-9_.])\$?([A-Z]{1,2})\$?(\d+)"
           r"(?::\$?([A-Z]{1,2})\$?(\d+))?(?![A-Za-z0-9_(])")

    def ev(sheet, f):
        def q(m):
            sh = m.group(1) or m.group(2) or sheet
            if m.group(5):
                return f"_R({sh!r},{m.group(3)!r},{m.group(4)!r},{m.group(5)!r},{m.group(6)!r})"
            return f"_C({sh!r},{m.group(3) + m.group(4)!r})"
        parts = f.split('"')
        expr = '"'.join(re.sub(pat, q, p) if i % 2 == 0 else p for i, p in enumerate(parts))
        return eval(expr, {"__builtins__": {}}, dict(ns, _R=rng, _C=cell))

    n = 0
    for sh in wb.sheetnames:
        for row in wb[sh].iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    v = cell(sh, c.coordinate); n += 1
                    assert v is not None and math.isfinite(v), (sh, c.coordinate, c.value, v)
    return cell, n


def main():
    runs = load()
    with PdfPages(OUT_PDF) as pdf:
        table_page(pdf, runs)
        for k, x in enumerate(runs, 2):
            run_page(pdf, x, k, len(runs) + 1)
    overview_png(runs)
    write_xlsx(runs)
    cell, n = evaluate_workbook(OUT_XLSX)
    for x in runs:
        rr = x["row"]
        checks = [(f"K{rr}", x["KD2_nM"])] + ([(f"F{rr}", x["ka1"]), (f"H{rr}", x["KD1_nM"])] if x["ka1_ok"] else [])
        for ref, want in checks:
            got = cell("Constants", ref)
            assert abs(got - want) <= 1e-7 * abs(want), (ref, got, want)
    for i, key in enumerate(("ka1", "kdis1", "KD1_nM", "ka2", "kdis2", "KD2_nM")):
        wt = [x[key] for x in runs if x["analyte"] == "EphB2-WT"]; mu = [x[key] for x in runs if x["analyte"] == "EphB2-R79P"]
        p = stats.ttest_ind(wt, mu, equal_var=True).pvalue
        assert abs(cell("WT vs R79P", f"I{i + 2}") - p) <= 1e-6 * p
    print(f"workbook: {n} formulas evaluated, all finite and matching Python")
    print(f"{'GF':3s} {'run':9s} {'analyte':11s} {'Ka1 1/Ms':>10s} {'Kdis1':>10s} {'KD1 nM':>7s} {'Ka2 1/Ms':>10s} {'Kdis2':>10s} "
          f"{'KD2 nM':>7s} {'R2 top':>7s} {'R2 all':>7s} {'R2 kobs':>7s} {'R2 dis':>7s}")
    for x in runs:
        ka1s, kd1s = (f"{x['ka1']:10.3e}", f"{x['KD1_nM']:7.3g}") if x["ka1_ok"] else (f"{'n.d.':>10s}", f"{'n.d.':>7s}")
        print(f"{x['gf']:3d} {x['exp'][:9]:9s} {x['analyte']:11s} {ka1s} {x['kdis1']:10.3e} {kd1s} "
              f"{x['ka2']:10.3e} {x['kdis2']:10.3e} {x['KD2_nM']:7.3g} {x['r2_top']:7.4f} {x['r2_all']:7.4f} "
              f"{x['r2_kobs']:7.3f} {x['r2_dis']:7.4f}")
    for key in ("ka1", "kdis1", "KD1_nM", "ka2", "kdis2", "KD2_nM"):
        wt = [x[key] for x in runs if x["analyte"] == "EphB2-WT"]; mu = [x[key] for x in runs if x["analyte"] == "EphB2-R79P"]
        p = stats.ttest_ind(wt, mu, equal_var=True).pvalue
        print(f"  {key:7s} WT {np.mean(wt):.3g} ± {np.std(wt, ddof=1):.2g}   R79P {np.mean(mu):.3g} ± {np.std(mu, ddof=1):.2g}   "
              f"ratio {np.mean(mu) / np.mean(wt):.2f}   p = {p:.3f}")
    print(OUT_PDF); print(OUT_PNG); print(OUT_XLSX)


if __name__ == "__main__":
    main()
