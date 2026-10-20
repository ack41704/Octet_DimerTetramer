"""
Rate constants for the extended runs (14-16 h dissociation), v3: POINT VALUES for every constant (no bounds, no n.d.),
so the replicate sets can be tested. Association split as in v2: FAST phase = DIMER, SLOW phase = TETRAMER.
Numbering as in the master table: 1 = tetramer, 2 = dimer.

  Ka1 (1/Ms)    v1 definition, now in every run: the top-concentration association trace refit alone (1500 evenly
                spaced points, fast + slow exponential; identical to extended_constants_all16.py), slow phase = tetramer,
                Ka1 = (k_slow - Kdis1) / [EphB2]_top.   (v1 called it n.d. in 2026.90/.96/.97 because it treated the slow
                phase there as the dimer; v3 does not.)
  Ka2 (1/Ms)    slope of the fast-phase (dimer) rate vs [EphB2] from the v2 global association fit (assoc_two_phase.py),
                used when the series determines it: profile closed on both sides AND >= 3 sensors whose fast phase is
                >= 10 % of their signal. Otherwise (2026.94, .98, .99) Ka2 = (k_fast - Kdis2) / [EphB2]_top with k_fast from
                the same top-concentration fit as Ka1; flagged APPROXIMATE per run (it assumes intercept = Kdis2).
  Kdis2, Kdis1  global biexponential dissociation fit (unchanged). 2026.91 and 2026.97 show no tetramer decay in 15 h
                (their other sensors drift upward); Kdis1 there is the pipeline's detection limit, reported as the value
                and noted. Neither run is part of a replicate test.
  KD1, KD2      Kdis/Ka in nM.
WT vs R79P Student's t-tests: EphrinB3-Fc in PBST (n = 3 each), all six constants. Buffers are never pooled.
Plots are the v2 ones (association fit coloured by phase), plus stars in the rate panel marking the top-concentration
rates the Ka1 (and fallback Ka2) values come from.

Outputs (analysis folder): Extended_constants_all16_v3_9-30-26.xlsx, Extended_sensorgrams_all16_v3_9-30-26.pdf,
Extended_sensorgrams_all16_v3_overview.png.
"""
from paths import DATA_DIR, WORK_DIR  # noqa: E402,F401  (repo copy: portable data / work locations)
import os
import sys
import math

import numpy as np
from scipy import stats
from scipy.optimize import curve_fit
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from eb3_extended_constants import two, evaluate_workbook  # noqa: E402
import extended_constants_all16_v2 as V  # noqa: E402
from extended_constants_all16_v2 import (INK, INK2, GRID, DIMER, TETRAMER, LN2, EARLY_S, sensorgram,  # noqa: E402
                                         early_panel, rate_panel, grp, r2, PHASE_HANDLES, PHASE_LABELS)

W = WORK_DIR
OUT_XLSX = os.path.join(W, "Extended_constants_all16_v3_9-30-26.xlsx")
OUT_PDF = os.path.join(W, "Extended_sensorgrams_all16_v3_9-30-26.pdf")
OUT_PNG = os.path.join(W, "Extended_sensorgrams_all16_v3_overview.png")
MIN_FAST_SENSORS, MIN_FAST_SHARE = 3, 0.10     # Ka2 slope needs >= 3 sensors whose fast phase is >= 10 % of their signal
# the 16 blue master-table runs + runs added later (fitted into fit_cache_extra.pkl by build_extra_cache.py)
EXTRA_RUNS = [("2026.122", 82, "EphrinA5-Fc", "EphB2-His")]
BUFFER_ASSUMED = {"2026.122"}                   # buffer not stated in the file name or CSV: PBST assumed
RUN_LIST = V.RUNS + EXTRA_RUNS


# ------------------------------------------------------------------ analysis
def top_trace_fit(x):
    """v1's top-concentration refit (extended_constants_all16.load): same trace choice, points, start values, bounds"""
    r = x["e"]["r"]; tb = r["tb"]
    w = max([w for w in r["keep"] if w["conc_nM"] == x["top"]], key=lambda w: w["y"].max())
    m = w["t"] <= tb
    t, y = w["t"][m], w["y"][m]
    i = np.unique(np.linspace(0, len(t) - 1, 1500).astype(int)); t, y = t[i], y[i]
    p, _ = curve_fit(two, t, y, p0=[y[-1] * .75, 5e-3, y[-1] * .25, 4e-4],
                     bounds=([0, 1e-4, 0, 1e-6], [5, 1, 5, 5e-3]), maxfev=40000)
    A1, k1, A2, k2 = p
    if k1 < k2:
        A1, k1, A2, k2 = A2, k2, A1, k1
    return dict(k_fast=float(k1), k_slow=float(k2), A_fast=float(A1), A_slow=float(A2),
                r2=float(1 - np.sum((y - two(t, A1, k1, A2, k2)) ** 2) / np.sum((y - y.mean()) ** 2)))


def load():
    return [finish_run(x) for x in V.load(RUN_LIST)]   # v2: global association fit, profiles, per-sensor rates


def finish_run(x, strict=True):
    """v3 point values for one v2-analysed run. strict: every value must be positive (the 17 extended runs).
    Non-strict (runs of any length): a value that cannot be formed (top-trace fit fails, Ka <= 0, Kdis not determinable)
    is NaN, shown as '—', and x["na_reasons"] says why."""
    C = x["top"] * 1e-9
    na = x.setdefault("na_reasons", {})
    try:
        tt = top_trace_fit(x)
    except Exception:                                  # non-strict only: no two-exponential fit of the top trace
        if strict:
            raise
        tt = None; na["ka1"] = "top-concentration trace could not be fitted with two exponentials"
    x["tt"] = tt
    x["ka1"] = (tt["k_slow"] - x["kdis1"]) / C if tt else np.nan            # v1 definition, slow phase = tetramer
    x["ka1_slope"] = x["fit"]["Ka1"]
    x["ka2_slope"] = x["fit"]["Ka2"]
    x["ka2_top"] = (tt["k_fast"] - x["kdis2"]) / C if tt else np.nan        # same top-trace fit as Ka1
    # the slope is used only if the concentration series determines it (profile closed on both sides) AND at least
    # MIN_FAST_SENSORS sensors carry a real fast phase; otherwise the slope rests on sensors with (almost) no dimer phase
    x["n_fast_sensors"] = sum(p["share_f"] >= MIN_FAST_SHARE for p in x["alone"])
    x["ka2_from_top"] = not x["prof2"]["resolved"] or x["n_fast_sensors"] < MIN_FAST_SENSORS
    x["ka2"] = x["ka2_top"] if x["ka2_from_top"] else x["ka2_slope"]
    if strict:
        assert x["ka1"] > 0 and x["ka2"] > 0, x["exp"]
    for k, why in (("ka1", "slow association rate at the top concentration is not faster than Kdis1"),
                   ("ka2", "fast association rate at the top concentration is not faster than Kdis2")):
        if not (np.isfinite(x[k]) and x[k] > 0):
            na.setdefault(k, why); x[k] = np.nan
    x["KD1_nM"] = x["kdis1"] / x["ka1"] * 1e9
    x["KD2_nM"] = x["kdis2"] / x["ka2"] * 1e9
    x["r2_top"] = tt["r2"] if tt else np.nan
    return x


def fv(v, f="{:.3e}"):
    """number as text, '—' when not determinable"""
    return f.format(v) if v is not None and np.isfinite(v) else "—"


def thalf(k, unit):
    return fv(LN2 / k / (86400 if unit == "d" else 60) if k and np.isfinite(k) else np.nan, "{:.2f}" if unit == "d"
              else "{:.1f}")


def ka2_reason(x):
    if not x["prof2"]["resolved"]:
        return "the fast-phase rate does not rise with [EphB2] across the series"
    return (f"only {x['n_fast_sensors']} sensor(s) carry a fast phase >= {100 * MIN_FAST_SHARE:.0f}% of their signal, so the "
            "slope is not determined by the series")


def notes(x):
    n = []
    if x["exp"] in BUFFER_ASSUMED:
        n.append(f"added 2026-09-30, not in the master table; buffer not stated in the file, {x['buffer']} assumed")
    if x["ka2_from_top"]:
        n.append(f"Ka2 APPROXIMATE, from the top concentration: {ka2_reason(x)} (global slope {x['ka2_slope']:.1e}); "
                 f"Ka2 = (k_fast - Kdis2)/[EphB2]top with k_fast = {x['tt']['k_fast']:.2e}/s from the fit used for Ka1, "
                 "which assumes the dimer k_obs intercept equals Kdis2")
    if x["kdis1_bound"]:
        n.append("no tetramer decay detected in 15 h (other sensors drift upward): Kdis1 = detection limit, so KD1 and "
                 "the tetramer half-life are limits")
    if x["kf0"] < 1e-6:
        n.append(f"dimer k_obs intercept ~0 (Kdis2 = {x['kdis2']:.1e}/s)")
    elif not 1 / 3 <= x["kf0"] / x["kdis2"] <= 3:
        n.append(f"dimer k_obs intercept ({x['kf0']:.1e}/s) differs from Kdis2 ({x['kdis2']:.1e}/s) by "
                 f"{max(x['kf0'] / x['kdis2'], x['kdis2'] / x['kf0']):.0f}x")
    if x["assoc_min"] < 60:
        n.append(f"{x['assoc_min']:.0f} min association: Ka1/KD1 not comparable with the 90-min runs (fitting only the first "
                 "20 min of the 90-min runs changes their slow rate 0.02-5x, 3-5x higher in the other EphrinB2 PBSTB runs)")
    return "; ".join(n)


# ------------------------------------------------------------------ figures
def rate_panel_v3(ax, x):
    rate_panel(ax, x)
    if x["tt"] is None:
        return
    ax.plot(x["top"], x["tt"]["k_slow"], "*", ms=10, mfc=TETRAMER, mec="white", mew=0.6, zorder=6)
    if x["ka2_from_top"]:
        ax.plot(x["top"], x["tt"]["k_fast"], "*", ms=10, mfc=DIMER, mec="white", mew=0.6, zorder=6)
    ax.text(0.98, 0.16, "★ = rate used for Ka1" + (" / Ka2" if x["ka2_from_top"] else ""), transform=ax.transAxes,
            fontsize=5.6, color=INK2, ha="right", va="bottom")


def run_page(pdf, x, k, n):
    fig = plt.figure(figsize=(11, 8.5))
    fig.text(0.05, 0.955, f"{x['lig']} + {x['analyte']}   —   {x['exp']}   (Global Fit # {x['gf']})",
             fontsize=13, weight="bold", color=INK)
    buf = f"{x['buffer']} (assumed: not stated in the file)" if x["exp"] in BUFFER_ASSUMED else x["buffer"]
    fig.text(0.05, 0.928, f"{buf} · association {x['assoc_min']:.0f} min · dissociation {x['dis_h']:.1f} h · "
             f"{x['n_fit']} of {x['n_all']} sensors fitted · dots = data by concentration (grey = excluded sensors)",
             fontsize=9, color=INK2)
    fig.text(0.95, 0.955, f"page {k} / {n}", fontsize=8, color="#777", ha="right")
    fig.legend(PHASE_HANDLES, PHASE_LABELS, loc="upper right", bbox_to_anchor=(0.975, 0.925), fontsize=7.5,
               frameon=False, title="association fit line", title_fontsize=7.5, alignment="left")
    ax = fig.add_axes([0.07, 0.4, 0.56, 0.49]); sensorgram(ax, x)
    ae = fig.add_axes([0.71, 0.63, 0.26, 0.2]); early_panel(ae, x, small=True)
    ae.set_title(f"association, first {EARLY_S // 60} min (dimer phase)", fontsize=8.5, color=INK2, loc="left")
    az = fig.add_axes([0.71, 0.4, 0.26, 0.17]); sensorgram(az, x, small=True, legend=False)
    az.set_xlim(-0.05, x["tb"] / 3600 + 1.0)
    az.set_title("association + first hour of dissociation", fontsize=8.5, color=INK2, loc="left")
    tt = x["tt"]; top = f"{x['top']:g} nM"; na = x.get("na_reasons", {})
    ka2_note = (na["ka2"] if "ka2" in na else
                f"APPROXIMATE: (k_fast − Kdis2)/[EphB2]top, k_fast = {tt['k_fast']:.2e}/s (★)" if x["ka2_from_top"]
                else "slope of fast (dimer) rate vs [EphB2], all curves")
    ka1_note = (na["ka1"] if "ka1" in na else
                f"(k_slow − Kdis1)/[EphB2]top, k_slow = {tt['k_slow']:.2e}/s at {top} (★)")
    kdis1_note = x.get("kdis1_note") or (f"half-life {thalf(x['kdis1'], 'd')} days" +
                                         ("  (no decay in 15 h: detection limit)" if x["kdis1_bound"] else ""))
    kdis2_note = x.get("kdis2_note") or f"half-life {thalf(x['kdis2'], 'min')} min"
    rows = [("Ka1  tetramer (1/Ms)", fv(x["ka1"]), ka1_note),
            ("Kdis1 tetramer (1/s)", fv(x["kdis1"]), kdis1_note),
            ("KD1  tetramer (nM)", fv(x["KD1_nM"], "{:.3g}"), "Kdis1 / Ka1"),
            ("Ka2  dimer (1/Ms)", fv(x["ka2"]), ka2_note),
            ("Kdis2 dimer (1/s)", fv(x["kdis2"]), kdis2_note),
            ("KD2  dimer (nM)", fv(x["KD2_nM"], "{:.3g}"), "Kdis2 / Ka2"),
            (f"signal at {top}", f"{100 * x['fast_share_top']:.0f}% / {100 * (1 - x['fast_share_top']):.0f}%",
             f"dimer / tetramer share at end of association; dimer dominates first {x['tx_top']:.0f} s"),
            ("R² association, top conc", fv(x["r2_top"], "{:.4f}"),
             "top-concentration fit used for Ka1" + (" and Ka2" if x["ka2_from_top"] else "")),
            ("R² association, all curves", f"{x['r2_all']:.4f}",
             ("global fit (phase colours)" if x["ka2_from_top"] else "global fit (Ka2)") +
             f"; first {EARLY_S // 60} min {x['r2_early']:.4f}; worst sensor {x['r2_min']:.4f}"),
            ("R² dissociation", f"{x['r2_dis']:.4f}", "global fit, all curves (Kdis1, Kdis2)")]
    if x["assoc_min"] < 60:
        rows[2] = (rows[2][0], rows[2][1], rows[2][2] + f"   ({x['assoc_min']:.0f}-min association: Ka1, KD1 not comparable "
                   "with the 90-min runs)")
    y = 0.315
    for name, val, note in rows:
        fig.text(0.05, y, name, fontsize=8.6, color=INK, family="DejaVu Sans Mono")
        fig.text(0.265, y, val, fontsize=8.6, color=INK, family="DejaVu Sans Mono", weight="bold")
        fig.text(0.39, y, note, fontsize=7.4, color=INK2)
        y -= 0.027
    ak = fig.add_axes([0.815, 0.075, 0.16, 0.225]); rate_panel_v3(ak, x)
    ak.set_title("rate vs [EphB2]", fontsize=8.5, color=INK2, loc="left")
    pdf.savefig(fig); plt.close(fig)


def table_page(pdf, runs):
    fig = plt.figure(figsize=(11, 8.5))
    fig.text(0.03, 0.955, "Extended runs (14-16 h dissociation): rate constants and fit quality — v3 (fast = dimer, "
             "slow = tetramer)", fontsize=13, weight="bold", color=INK)
    fig.text(0.03, 0.93, "1 = tetramer, 2 = dimer. Ka1 = (k_slow − Kdis1)/[EphB2]top from the top-concentration slow "
             "phase (as in v1); Ka2 = slope of the fast-phase rate vs [EphB2], all curves (global fit);\nKdis from the "
             "global dissociation fit. Point values throughout.", fontsize=7.8, color=INK2, va="top", linespacing=1.5)
    ax = fig.add_axes([0.01, 0.1, 0.98, 0.8]); ax.axis("off")
    head = ["GF", "Run", "Ligand", "Analyte", "Buffer", "Ka1\n(1/Ms)", "Kdis1\n(1/s)", "KD1\n(nM)", "Ka2\n(1/Ms)",
            "Kdis2\n(1/s)", "KD2\n(nM)", "R² assoc\ntop", "R² assoc\nall", "R²\ndissoc"]
    rows = [[str(x["gf"]), x["exp"].split()[0], x["lig"].replace("Ephrin", "E").replace("-Fc", ""), x["analyte"],
             x["buffer"], f"{x['ka1']:.3e}", f"{x['kdis1']:.3e}", f"{x['KD1_nM']:.3g}", f"{x['ka2']:.3e}",
             f"{x['kdis2']:.3e}", f"{x['KD2_nM']:.3g}", f"{x['r2_top']:.4f}", f"{x['r2_all']:.4f}",
             f"{x['r2_dis']:.4f}"] for x in runs]
    wd = [0.035, 0.065, 0.045, 0.085, 0.055, 0.085, 0.085, 0.06, 0.085, 0.085, 0.055, 0.065, 0.065, 0.06]
    tb = ax.table(cellText=rows, colLabels=head, loc="upper center", cellLoc="center", colWidths=[w / sum(wd) for w in wd])
    tb.auto_set_font_size(False); tb.set_fontsize(7.4); tb.scale(1, 1.62)
    for (i, j), c in tb.get_celld().items():
        c.set_edgecolor(GRID)
        if i == 0:
            c.set_facecolor("#f2f2ef"); c.set_text_props(weight="bold", color=INK)
    top2 = ", ".join(x["exp"].split()[0] for x in runs if x["ka2_from_top"])
    lim = ", ".join(x["exp"].split()[0] for x in runs if x["kdis1_bound"])
    fig.text(0.03, 0.03, f"Ka2 / KD2 APPROXIMATE in {top2}: the concentration series does not determine the fast-phase slope, "
             "so Ka2 = (k_fast − Kdis2)/[EphB2]top from the top-concentration fit.\n"
             f"Kdis1 in {lim}: no tetramer decay detected in 15 h; the value is the detection limit (KD1 and the tetramer "
             "half-life are limits).\n2026.90: 20-min association, so its Ka1/KD1 are not comparable with the 90-min runs.  "
             "None of these runs is in the WT vs R79P t-tests.\n"
             + "".join(f"{e} (GF {g}, {l}) added 9-30-26: buffer not stated in the file, PBST assumed.\n"
                       for e, g, l, _ in EXTRA_RUNS if e in BUFFER_ASSUMED) + "Ka2/KD2 are fast-phase (dimer) values, so they differ from "
             "v1 in every run (PBST WT vs R79P: v1 Ka2 p = 0.646, KD2 p = 0.010; v3 0.876, 0.015).",
             fontsize=7.2, color=INK2, linespacing=1.5)
    pdf.savefig(fig); plt.close(fig)


KEYS = (("ka1", "Ka1 (1/Ms)", "{:.3e}"), ("kdis1", "Kdis1 (1/s)", "{:.3e}"), ("KD1_nM", "KD1 (nM)", "{:.3g}"),
        ("ka2", "Ka2 (1/Ms)", "{:.3e}"), ("kdis2", "Kdis2 (1/s)", "{:.3e}"), ("KD2_nM", "KD2 (nM)", "{:.3g}"))


def ttests(runs):
    wt, mu = grp(runs, "EphB2-WT"), grp(runs, "EphB2-R79P")
    assert len(wt) == len(mu) == 3
    return [dict(key=k, name=nm, f=f, a=[x[k] for x in wt], b=[x[k] for x in mu],
                 p=stats.ttest_ind([x[k] for x in wt], [x[k] for x in mu], equal_var=True).pvalue) for k, nm, f in KEYS]


def stats_page(pdf, runs):
    fig = plt.figure(figsize=(11, 8.5))
    fig.text(0.04, 0.95, "EphB2-WT vs EphB2-R79P", fontsize=13, weight="bold", color=INK)
    fig.text(0.04, 0.935, "Student's t-test, two-tailed, equal variance. Only EphrinB3-Fc in PBST has replicates (n = 3 "
             "each);\nthe PBSTB sets have one run per variant, shown below for reference (not tested). Buffers are not "
             "pooled.", fontsize=8.3, color=INK2, va="top", linespacing=1.5)
    ax2 = fig.add_axes([0.03, 0.5, 0.94, 0.38]); ax2.axis("off")
    rows2 = [[t["name"], f"{t['f'].format(np.mean(t['a']))} ± {t['f'].format(np.std(t['a'], ddof=1))}",
              f"{t['f'].format(np.mean(t['b']))} ± {t['f'].format(np.std(t['b'], ddof=1))}",
              f"{np.mean(t['b']) / np.mean(t['a']):.2f}", "<0.001" if t["p"] < 0.001 else f"{t['p']:.3f}"]
             for t in ttests(runs)]
    tb2 = ax2.table(cellText=rows2, colLabels=["EphrinB3-Fc, PBST", "WT mean ± SD (n=3)", "R79P mean ± SD (n=3)",
                                               "R79P / WT", "Student's p"], loc="upper center", cellLoc="center")
    tb2.auto_set_font_size(False); tb2.set_fontsize(8.4); tb2.scale(1, 1.8)
    ax3 = fig.add_axes([0.03, 0.08, 0.94, 0.36]); ax3.axis("off")
    rows3 = []
    for lig in ("EphrinB3-Fc", "EphrinB2-Fc"):
        a, b = grp(runs, "EphB2-WT", lig, "PBSTB")[0], grp(runs, "EphB2-R79P", lig, "PBSTB")[0]
        for key, name, f in KEYS:
            rows3.append([f"{lig}, PBSTB", name, f"{f.format(a[key])} ({a['exp'].split()[0]})",
                          f"{f.format(b[key])} ({b['exp'].split()[0]})"])
    tb3 = ax3.table(cellText=rows3, colLabels=["Set", "Constant", "WT (single run)", "R79P (single run)"],
                    loc="upper center", cellLoc="center")
    tb3.auto_set_font_size(False); tb3.set_fontsize(7.6); tb3.scale(1, 1.3)
    for t in (tb2, tb3):
        for (i, j), c in t.get_celld().items():
            c.set_edgecolor(GRID)
            if i == 0:
                c.set_facecolor("#f2f2ef"); c.set_text_props(weight="bold", color=INK)
    pdf.savefig(fig); plt.close(fig)


# ------------------------------------------------------------------ workbook
COLS = ["Global Fit #", "Run", "Ligand", "Analyte", "Buffer", "Association (min)", "Dissociation (h)",
        "Top concentration (nM)",
        "Ka1 tetramer (1/Ms)", "Slow (tetramer) rate at top conc, k_slow (1/s)", "Kdis1 tetramer (1/s)", "Kdis1 source",
        "KD1 tetramer (nM)",
        "Ka2 dimer (1/Ms)", "Ka2 source", "Ka2: slope of fast-phase rate vs [EphB2], all curves (1/Ms)",
        "Ka2: top conc = (k_fast - Kdis2)/[EphB2]top (1/Ms)", "Fast (dimer) rate at top conc, k_fast (1/s)",
        "Kdis2 dimer (1/s)", "KD2 dimer (nM)",
        "R² association, top concentration (Ka1)", "R² association, all curves (Ka2)",
        "R² association, first 10 min, all curves", "R² dissociation, global fit (Kdis1, Kdis2)",
        "Tetramer half-life (days)", "Dimer half-life (min)", "Notes"]


def write_xlsx(runs):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter as gl
    wb = Workbook()
    F, FB = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
    FILL = PatternFill("solid", fgColor="F2F2EF"); BOX = Border(bottom=Side(style="thin", color="D9D9D6"))
    col = {name: j for j, name in enumerate(COLS, 1)}
    L = {name: gl(j) for name, j in col.items()}

    def header(ws, cols, widths, freeze="E2", height=60):
        for j, (h, w) in enumerate(zip(cols, widths), 1):
            c = ws.cell(row=1, column=j, value=h); c.font = FB; c.fill = FILL; c.border = BOX
            c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
            ws.column_dimensions[c.column_letter].width = w
        ws.row_dimensions[1].height = height; ws.freeze_panes = freeze

    ws = wb.active; ws.title = "Constants"
    header(ws, COLS, [8, 15, 12, 12, 8, 10, 10, 10, 13, 14, 13, 16, 11, 13, 14, 15, 15, 13, 13, 10, 12, 12, 12, 13, 11,
                      11, 80])
    num = lambda v: float(f"{v:.10g}")
    E3 = "0.000E+00"
    fmt = {"Ka1 tetramer (1/Ms)": E3, "Slow (tetramer) rate at top conc, k_slow (1/s)": E3, "Kdis1 tetramer (1/s)": E3,
           "KD1 tetramer (nM)": "0.0000", "Ka2 dimer (1/Ms)": E3, COLS[15]: E3, COLS[16]: E3,
           "Fast (dimer) rate at top conc, k_fast (1/s)": E3, "Kdis2 dimer (1/s)": E3, "KD2 dimer (nM)": "0.00",
           "R² association, top concentration (Ka1)": "0.0000", "R² association, all curves (Ka2)": "0.0000",
           "R² association, first 10 min, all curves": "0.0000", "R² dissociation, global fit (Kdis1, Kdis2)": "0.0000",
           "Tetramer half-life (days)": "0.00", "Dimer half-life (min)": "0.0", "Association (min)": "0",
           "Dissociation (h)": "0.0", "Top concentration (nM)": "0.0"}
    for i, x in enumerate(runs):
        rr = i + 2; x["row"] = rr
        c_ = lambda name: f"{L[name]}{rr}"
        v = {"Global Fit #": x["gf"], "Run": x["exp"], "Ligand": x["lig"], "Analyte": x["analyte"], "Buffer": x["buffer"],
             "Association (min)": round(x["assoc_min"], 1), "Dissociation (h)": round(x["dis_h"], 2),
             "Top concentration (nM)": x["top"],
             "Ka1 tetramer (1/Ms)": f"=({c_('Slow (tetramer) rate at top conc, k_slow (1/s)')}-{c_('Kdis1 tetramer (1/s)')})"
                                    f"/({c_('Top concentration (nM)')}*1E-9)",
             "Slow (tetramer) rate at top conc, k_slow (1/s)": num(x["tt"]["k_slow"]),
             "Kdis1 tetramer (1/s)": num(x["kdis1"]),
             "Kdis1 source": "detection limit (no decay in 15 h)" if x["kdis1_bound"] else "global dissociation fit",
             "KD1 tetramer (nM)": f"={c_('Kdis1 tetramer (1/s)')}/{c_('Ka1 tetramer (1/Ms)')}*1E9",
             "Ka2 dimer (1/Ms)": f"={c_(COLS[16])}" if x["ka2_from_top"] else f"={c_(COLS[15])}",
             "Ka2 source": "top concentration, APPROXIMATE (slope not determined by the series)" if x["ka2_from_top"]
             else "slope, all curves",
             COLS[15]: num(x["ka2_slope"]),
             COLS[16]: f"=({c_('Fast (dimer) rate at top conc, k_fast (1/s)')}-{c_('Kdis2 dimer (1/s)')})"
                       f"/({c_('Top concentration (nM)')}*1E-9)",
             "Fast (dimer) rate at top conc, k_fast (1/s)": num(x["tt"]["k_fast"]),
             "Kdis2 dimer (1/s)": num(x["kdis2"]),
             "KD2 dimer (nM)": f"={c_('Kdis2 dimer (1/s)')}/{c_('Ka2 dimer (1/Ms)')}*1E9",
             "R² association, top concentration (Ka1)": num(x["r2_top"]),
             "R² association, all curves (Ka2)": num(x["r2_all"]),
             "R² association, first 10 min, all curves": num(x["r2_early"]),
             "R² dissociation, global fit (Kdis1, Kdis2)": num(x["r2_dis"]),
             "Tetramer half-life (days)": f"=LN(2)/{c_('Kdis1 tetramer (1/s)')}/86400",
             "Dimer half-life (min)": f"=LN(2)/{c_('Kdis2 dimer (1/s)')}/60", "Notes": notes(x)}
        for name, val in v.items():
            c = ws.cell(row=rr, column=col[name], value=val); c.font = F
            if name in fmt:
                c.number_format = fmt[name]
        ws.cell(row=rr, column=col["Notes"]).alignment = Alignment(wrap_text=True, vertical="top")

    st = wb.create_sheet("WT vs R79P")
    header(st, ["Constant (EphrinB3-Fc, PBST)", "WT mean", "WT SD", "WT n", "R79P mean", "R79P SD", "R79P n", "R79P / WT",
                "Student's t-test p (two-tailed, equal variance)"], [28, 13, 13, 7, 13, 13, 8, 10, 18], freeze="A2",
           height=46)
    wr = [x["row"] for x in grp(runs, "EphB2-WT")]; mr = [x["row"] for x in grp(runs, "EphB2-R79P")]
    assert wr == list(range(wr[0], wr[-1] + 1)) and mr == list(range(mr[0], mr[-1] + 1)) and len(wr) == len(mr) == 3
    cols6 = {"ka1": "Ka1 tetramer (1/Ms)", "kdis1": "Kdis1 tetramer (1/s)", "KD1_nM": "KD1 tetramer (nM)",
             "ka2": "Ka2 dimer (1/Ms)", "kdis2": "Kdis2 dimer (1/s)", "KD2_nM": "KD2 dimer (nM)"}
    for i, (key, name, _) in enumerate(KEYS):
        rr = i + 2; c = L[cols6[key]]
        a, b = f"Constants!{c}{wr[0]}:{c}{wr[-1]}", f"Constants!{c}{mr[0]}:{c}{mr[-1]}"
        for j, val in enumerate([name, f"=AVERAGE({a})", f"=STDEV({a})", f"=COUNT({a})", f"=AVERAGE({b})",
                                 f"=STDEV({b})", f"=COUNT({b})", f"=E{rr}/B{rr}", f"=TTEST({a},{b},2,2)"], 1):
            st.cell(row=rr, column=j, value=val).font = F
        for cc in (2, 3, 5, 6):
            st.cell(row=rr, column=cc).number_format = fmt[cols6[key]]
        st.cell(row=rr, column=8).number_format = "0.00"; st.cell(row=rr, column=9).number_format = "0.000"
    st.cell(row=9, column=1, value=("WT = GF 69, 70, 73; R79P = GF 71, 72, 74. The PBSTB sets (EphrinB3: GF 67 vs 68; "
                                    "EphrinB2: GF 60 vs 61) have one run per variant and are not tested. Buffers are not pooled.")
            ).font = Font(name="Arial", size=9, italic=True)

    dg = wb.create_sheet("Fit diagnostics")
    dcols = ["Global Fit #", "Run", "Ka1 slope from the global fit (1/Ms), for comparison", "Tetramer k_obs intercept ks0 (1/s)",
             "Dimer k_obs intercept kf0 (1/s)", "Kdis2 (1/s)", "Dimer share of association signal at top conc",
             "Fast-off share of dissociation amplitude at top conc", "Dimer-dominated time at top conc (s)",
             "Top-conc fit: fast rate (1/s)", "Top-conc fit: slow rate (1/s)", "R² association, worst sensor",
             "Ka2 slope 5% χ² profile closed on both sides?", "Sensors with a fast phase >= 10% of their signal"]
    header(dg, dcols, [8, 15, 16, 14, 14, 12, 14, 14, 13, 13, 13, 12, 14, 14], freeze="C2", height=72)
    for i, x in enumerate(runs):
        rr = i + 2
        vals = [x["gf"], x["exp"], num(x["ka1_slope"]), num(x["ks0"]), num(x["kf0"]), num(x["kdis2"]),
                num(x["fast_share_top"]), num(x["fastoff_share_top"]), round(x["tx_top"], 1), num(x["tt"]["k_fast"]),
                num(x["tt"]["k_slow"]), num(x["r2_min"]), "yes" if x["prof2"]["resolved"] else "no", x["n_fast_sensors"]]
        for j, (val, f) in enumerate(zip(vals, [None, None, E3, E3, E3, E3, "0%", "0%", "0", E3, E3, "0.0000", None,
                                                None]), 1):
            c = dg.cell(row=rr, column=j, value=val); c.font = F
            if f:
                c.number_format = f

    nt = wb.create_sheet("Notes"); nt.column_dimensions["A"].width = 130
    for i, t in enumerate([
        "v3 of the extended runs (14-16 h dissociation): the 16 blue-highlighted rows of 'MASTER TABLE (BEST Colored)' plus "
        "2026.122 (EphrinA5-Fc + EphB2-His, Global Fit 82; added 2026-09-30, not in the master table, buffer not stated in "
        "the file so PBST assumed). Generated 2026-09-30. 1 = tetramer, 2 = dimer. Point values for every constant "
        "(no bounds, no n.d.).",
        "Association split (as in v2): FAST phase = dimer, SLOW phase = tetramer, in every run and at every concentration.",
        "Ka1 (column I, formula): v1 definition, now in every run. The top-concentration association trace is refit alone "
        "(1500 evenly spaced points, fast + slow exponential, as in v1) and its slow rate is the tetramer rate (column J): "
        "Ka1 = (k_slow - Kdis1)/[EphB2]_top. This reproduces v1's Ka1 in the 13 runs v1 reported; v1's n.d. runs "
        "(2026.90, .96, .97) now get a value because their slow phase is taken as the tetramer.",
        "Ka2 (column N, formula): slope of the fast-phase (dimer) rate vs [EphB2] from one global association fit per run "
        "(fast rate = Ka2*C + kf0, slow rate = Ka1*C + ks0, amplitudes per sensor; column P), used when the series "
        "determines it: 5% chi-square profile closed on both sides AND >= 3 sensors whose fast phase is >= 10% of their "
        "signal. Otherwise (2026.94, .98, .99; see 'Fit diagnostics' M-N) Ka2 = (k_fast - Kdis2)/[EphB2]_top (column Q) "
        "with k_fast from the same top-concentration fit as Ka1 (column R). These three are APPROXIMATE: the formula "
        "assumes the dimer k_obs intercept equals Kdis2, which the global fit does not support there. Column O says which.",
        "Ka2 and KD2 differ from v1 in every run: v1 took the per-sensor component with the larger amplitude (often the "
        "tetramer phase), v3 the fast (dimer) phase. PBST WT vs R79P: v1 Ka2 p = 0.646, KD2 p = 0.010; v3 0.876, 0.015. "
        "Ka1, Kdis1, KD1 and Kdis2 are unchanged from v1 wherever v1 gave a number.",
        "2026.90 has a 20-min association: its Ka1 and KD1 are not comparable with the 90-min runs (fitting only the first "
        "20 min of the other runs changes their top-concentration slow rate 0.02-5x, and 3-5x upward in the other EphrinB2 "
        "PBSTB runs).",
        "Kdis2, Kdis1: fast and slow rates of the global biexponential dissociation fit (unchanged). In 2026.91 and 2026.97 "
        "no tetramer decay is detected in 15 h (most sensors drift upward), so Kdis1 is the pipeline's detection limit, "
        "used as the value (column L). KD1 and the tetramer half-life of those two runs are therefore limits. Neither run "
        "is in a replicate test.",
        "KD = Kdis/Ka x 1e9 (nM). Half-lives = LN(2)/Kdis.",
        "R2 association (top concentration) = the top-concentration fit used for Ka1; R2 association (all curves) = the "
        "global fit used for Ka2, on the full 0.5 s data; R2 dissociation = pooled global dissociation fit.",
        "'Fit diagnostics' sheet: the global-fit Ka1 slope (not used: its slow-rate intercept ks0 is 100-1000x Kdis1, so the "
        "slow phase is largely concentration-independent), the k_obs intercepts vs Kdis2, phase shares, and the two "
        "criteria that decide the Ka2 source.",
        "WT vs R79P: EphrinB3-Fc in PBST, n = 3 each (GF 69, 70, 73 vs 71, 72, 74), Student's t-test two-tailed equal "
        "variance, as formulas on the Constants sheet.",
        "Typed values are fit outputs; Ka1, Ka2, KD1, KD2, the top-concentration Ka2, the half-lives and all statistics "
        "are formulas.",
        "Code: dimer_tetramer/extended_constants_all16_v3.py "
        "(+ extended_constants_all16_v2.py, assoc_two_phase.py)"], 1):
        c = nt.cell(row=i, column=1, value=t); c.font = F; c.alignment = Alignment(wrap_text=True)
    wb.save(OUT_XLSX)
    return L


def main():
    runs = load()
    with PdfPages(OUT_PDF) as pdf:
        table_page(pdf, runs)
        stats_page(pdf, runs)
        for k, x in enumerate(runs, 3):
            run_page(pdf, x, k, len(runs) + 2)
        pdf.infodict()["Title"] = f"All {len(runs)} extended runs v3: point values, fast phase = dimer, slow phase = tetramer"
    V.overview_png(runs, path=OUT_PNG, tag="v3")
    L = write_xlsx(runs)
    cell, n = evaluate_workbook(OUT_XLSX)
    for x in runs:
        rr = x["row"]
        for name, want in (("Ka1 tetramer (1/Ms)", x["ka1"]), ("KD1 tetramer (nM)", x["KD1_nM"]),
                           ("Ka2 dimer (1/Ms)", x["ka2"]), ("KD2 dimer (nM)", x["KD2_nM"]),
                           (COLS[16], x["ka2_top"])):
            got = cell("Constants", f"{L[name]}{rr}")
            assert abs(got - want) <= 1e-7 * abs(want), (x["exp"], name, got, want)
    for i, t in enumerate(ttests(runs)):
        got = cell("WT vs R79P", f"I{i + 2}"); assert abs(got - t["p"]) <= 1e-6 * t["p"], (t["key"], got, t["p"])
    print(f"workbook: {n} formulas evaluated, all finite and matching Python")
    print(f"{'GF':3s} {'run':9s} {'lig':4s} {'analyte':11s} {'buf':6s} {'Ka1':>10s} {'Kdis1':>10s} {'KD1 nM':>8s} "
          f"{'Ka2':>10s} {'Kdis2':>10s} {'KD2 nM':>7s} {'R2top':>6s} {'R2all':>6s} {'R2dis':>6s}  Ka2 src")
    for x in runs:
        print(f"{x['gf']:3d} {x['exp'].split()[0]:9s} {x['lig'][6:8]:4s} {x['analyte']:11s} {x['buffer']:6s} "
              f"{x['ka1']:10.3e} {x['kdis1']:10.3e} {x['KD1_nM']:8.3g} {x['ka2']:10.3e} {x['kdis2']:10.3e} "
              f"{x['KD2_nM']:7.3g} {x['r2_top']:6.4f} {x['r2_all']:6.4f} {x['r2_dis']:6.4f}  "
              f"{'top' if x['ka2_from_top'] else 'slope'}{'  Kdis1=det.limit' if x['kdis1_bound'] else ''}")
    for t in ttests(runs):
        print(f"  EB3 PBST {t['key']:7s} WT {np.mean(t['a']):.4g} ± {np.std(t['a'], ddof=1):.3g}   "
              f"R79P {np.mean(t['b']):.4g} ± {np.std(t['b'], ddof=1):.3g}   ratio {np.mean(t['b']) / np.mean(t['a']):.2f}   "
              f"p = {t['p']:.4f}")
    print(OUT_PDF); print(OUT_PNG); print(OUT_XLSX)


if __name__ == "__main__":
    main()
