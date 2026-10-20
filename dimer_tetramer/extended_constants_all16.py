"""
Rate constants for ALL 16 extended runs (14-16 h dissociation) -- the blue-highlighted rows of
"MASTER TABLE (BEST Colored)" in Global Fit - Master Table on 9-29-26.numbers. Same analysis as
eb3_extended_constants.py (whose plotting and workbook-verification helpers are reused).
Numbering as in the master table: 1 = tetramer, 2 = dimer.

  Ka2 (1/Ms)    slope of k_obs vs [EphB2] over all association curves (pipeline k_obs regression)
  Kdis2, Kdis1  fast / slow rate of the global biexponential dissociation fit (all fitted sensors).
                If no tetramer decay was detected in 15 h (2026.91, 2026.97), Kdis1 is the pipeline's
                upper bound and every value derived from it is a bound (KD1 <, tetramer t1/2 >).
  Ka1 (1/Ms)    top concentration only: its association trace refit alone (1500 evenly spaced points,
                fast + slow exponential); the slow phase is treated as a second, bimolecular step,
                k_slow = Ka1*[EphB2] + Kdis1  ->  Ka1 = (k_slow - Kdis1) / [EphB2]_top.
                Determinable only if the FAST component is the dimer binding, i.e. its rate is closer than
                the slow rate to the dimer k_obs expected at that concentration (Ka2*[EphB2] + Kdis2).
                Otherwise the "slow" component is the dimer itself (2026.90, .96, .97) and Ka1 is n.d.
  KD1, KD2      Kdis/Ka in nM.   R^2: association top conc (Ka1); association all curves pooled (Ka2);
                k_obs regression (Ka2); dissociation global fit (Kdis1, Kdis2)
WT vs R79P Student's t-tests: EphrinB3-Fc in PBST only (n = 3 each); every other ligand/buffer set has one
run per variant. Buffers are never pooled.

Outputs: Extended_constants_all16_9-30-26.xlsx, Extended_sensorgrams_all16_9-30-26.pdf,
Extended_sensorgrams_all16_overview.png (analysis folder).
"""
from paths import DATA_DIR, WORK_DIR  # noqa: E402,F401  (repo copy: portable data / work locations)
import os
import sys
import math
import pickle

import numpy as np
from scipy import stats
from scipy.optimize import curve_fit
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from batch_fit import f2a  # noqa: E402
from eb3_extended_constants import sensorgram, two, evaluate_workbook  # noqa: E402

W = WORK_DIR
OUT_XLSX = os.path.join(W, "Extended_constants_all16_9-30-26.xlsx")
OUT_PDF = os.path.join(W, "Extended_sensorgrams_all16_9-30-26.pdf")
OUT_PNG = os.path.join(W, "Extended_sensorgrams_all16_overview.png")
# (experiment as named in the CSVs, Global Fit #, ligand, analyte) -- the 16 blue rows of the BEST Colored tab,
# ordered EphrinB3 PBST, EphrinB3 PBSTB, EphrinB2 PBSTB; WT then R79P then His-tagged within each set
RUNS = [("2026.101", 69, "EphrinB3-Fc", "EphB2-WT"), ("2026.103", 70, "EphrinB3-Fc", "EphB2-WT"),
        ("2026.108", 73, "EphrinB3-Fc", "EphB2-WT"), ("2026.106", 71, "EphrinB3-Fc", "EphB2-R79P"),
        ("2026.107", 72, "EphrinB3-Fc", "EphB2-R79P"), ("2026.109", 74, "EphrinB3-Fc", "EphB2-R79P"),
        ("2026.111", 75, "EphrinB3-Fc", "EphB2-His"), ("2026.116", 77, "EphrinB3-Fc", "EphB1-His"),
        ("2026.98 (PBSTB)", 67, "EphrinB3-Fc", "EphB2-WT"), ("2026.99 (PBSTB)", 68, "EphrinB3-Fc", "EphB2-R79P"),
        ("2026.96 (PBSTB)", 65, "EphrinB3-Fc", "EphB2-His"), ("2026.97 (PBSTB)", 66, "EphrinB3-Fc", "EphB2-His"),
        ("2026.91 (PBSTB)", 60, "EphrinB2-Fc", "EphB2-WT"), ("2026.92 (PBSTB)", 61, "EphrinB2-Fc", "EphB2-R79P"),
        ("2026.90 (PBSTB)", 59, "EphrinB2-Fc", "EphB2-His"), ("2026.94 (PBSTB)", 64, "EphrinB2-Fc", "EphB2-His")]
INK, INK2, GRID = "#1f1f1f", "#5a5a58", "#e6e6e3"
LN2 = math.log(2)


def assoc_fit_per_well(r):
    """The pipeline's per-sensor association fit (plot_all_fits.stage2_full: same points, start values, bounds),
    kept per sensor. stage2_full stores fits by concentration, so of two sensors at one concentration only the
    last survives (2026.97, 27 nM). For every other sensor this reproduces the stored fit exactly (asserted)."""
    tb, out = r["tb"], []
    for w in r["keep"]:
        m = w["t"] <= tb
        ta, ya = w["t"][m], w["y"][m]
        if len(ta) < 40:
            out.append(None); continue
        i = np.unique(np.geomspace(1, len(ta) - 1, 300).astype(int))
        span = max(ya[i][-1], 1e-6)
        try:
            p, _ = curve_fit(f2a, ta[i], ya[i], p0=[span * .5, 5e-3, span * .5, 4e-4],
                             bounds=([0, 1e-5, 0, 1e-6], [3, 1, 3, 1e-1]), maxfev=40000)
        except Exception:
            out.append(None); continue
        A1, k1, A2, k2 = p
        if k1 < k2:
            A1, k1, A2, k2 = A2, k2, A1, k1
        out.append((A1, k1, A2, k2))
    concs = [w["conc_nM"] for w in r["keep"]]
    for w, pa in zip(r["keep"], out):
        if concs.count(w["conc_nM"]) == 1 and w["conc_nM"] in r["assoc_fits"]:
            assert np.allclose(pa, r["assoc_fits"][w["conc_nM"]], rtol=1e-9), (w["conc_nM"], pa)
    return out


def load():
    cache = pickle.load(open(os.path.join(W, "fit_cache.pkl"), "rb"))
    by = {(e["r"]["interaction"], e["r"]["experiment"]): e for e in cache.values() if "error" not in e}
    runs = []
    for exp, gf, lig, analyte in RUNS:
        e = by[(f"{lig} + {analyte}", exp)]; r, s = e["r"], e["s"]; tb = r["tb"]
        assert s["dissoc_h"] >= 14, exp
        top = max(w["conc_nM"] for w in r["keep"])
        w = max([w for w in r["keep"] if w["conc_nM"] == top], key=lambda w: w["y"].max())
        m = w["t"] <= tb
        t, y = w["t"][m], w["y"][m]
        i = np.unique(np.linspace(0, len(t) - 1, 1500).astype(int)); t, y = t[i], y[i]
        p, _ = curve_fit(two, t, y, p0=[y[-1] * .75, 5e-3, y[-1] * .25, 4e-4],
                         bounds=([0, 1e-4, 0, 1e-6], [5, 1, 5, 5e-3]), maxfev=40000)
        A1, k1, A2, k2 = p
        if k1 < k2:
            A1, k1, A2, k2 = A2, k2, A1, k1
        r2_top = 1 - np.sum((y - two(t, A1, k1, A2, k2)) ** 2) / np.sum((y - y.mean()) ** 2)
        per_well = assoc_fit_per_well(r)
        ss_res = 0.0; ally = []
        for ww, pa in zip(r["keep"], per_well):
            if pa is not None:
                ma = ww["t"] <= tb
                ya = ww["y"][ma]; ally.append(ya)
                ss_res += np.sum((ya - f2a(ww["t"][ma], *pa)) ** 2)
        ally = np.concatenate(ally)
        measured = s["tet_category"] == "measured"
        kdis1 = r["ks"] if measured else s["tet_kdis2_upper"]
        k_exp = r["ka1"] * top * 1e-9 + r["kf"]                 # dimer k_obs expected at the top concentration
        ka1_ok = abs(math.log10(k1 / k_exp)) < abs(math.log10(k2 / k_exp)) and k1 / k2 >= 3
        x = dict(exp=exp, gf=gf, lig=lig, analyte=analyte, e=e, buffer=s["buffer"], top=top, fast=(A1, k1),
                 assoc_by_well=per_well,
                 slow=(A2, k2), k_fast=k1, k_slow=k2, k_exp=k_exp, ka1_ok=ka1_ok, kdis1=kdis1,
                 kdis1_bound=not measured, kdis2=r["kf"], ka2=r["ka1"], r2_top=float(r2_top),
                 r2_all=float(1 - ss_res / np.sum((ally - ally.mean()) ** 2)), r2_kobs=r["r2k"], r2_dis=r["R2d"],
                 assoc_min=s["assoc_min"], dis_h=s["dissoc_h"], n_fit=len(r["keep"]), n_all=len(e["mask"]),
                 slow_share=A2 * (1 - np.exp(-k2 * tb)) / two(tb, A1, k1, A2, k2))
        x["ka1"] = (k2 - kdis1) / (top * 1e-9)
        x["KD1_nM"] = kdis1 / x["ka1"] * 1e9
        x["KD2_nM"] = r["kf"] / r["ka1"] * 1e9
        runs.append(x)
    return runs


def ka1_text(x, f="{:.3e}"):
    return f.format(x["ka1"]) if x["ka1_ok"] else "n.d."


def kd1_text(x):
    if not x["ka1_ok"]:
        return "n.d."
    return ("< " if x["kdis1_bound"] else "") + f"{x['KD1_nM']:.3g}"


def kdis1_text(x):
    return ("< " if x["kdis1_bound"] else "") + f"{x['kdis1']:.3e}"


def t1_text(x):
    return ("> " if x["kdis1_bound"] else "") + f"{LN2 / x['kdis1'] / 86400:.2f}"


def notes(x):
    n = []
    if not x["ka1_ok"]:
        n.append(f"Ka1 n.d.: at {x['top']:g} nM the slow component (k = {x['k_slow']:.2e}) matches the expected dimer "
                 f"k_obs ({x['k_exp']:.2e}), so it is the dimer binding, not a second phase")
    if x["kdis1_bound"]:
        n.append("no tetramer decay detected in 15 h: Kdis1 is an upper bound (KD1 <, tetramer t1/2 >)")
    if x["assoc_min"] < 60:
        n.append(f"{x['assoc_min']:.0f} min association")
    if x["r2_kobs"] < 0.9:
        n.append(f"Ka2 regression R2 {x['r2_kobs']:.2f} (<0.9): Ka2 and KD2 approximate")
    return "; ".join(n)


# ------------------------------------------------------------------ PDF
def run_page(pdf, x, k, n):
    fig = plt.figure(figsize=(11, 8.5))
    fig.text(0.05, 0.955, f"{x['lig']} + {x['analyte']}   —   {x['exp']}   (Global Fit # {x['gf']})",
             fontsize=13, weight="bold", color=INK)
    fig.text(0.05, 0.928, f"{x['buffer']} · association {x['assoc_min']:.0f} min · dissociation {x['dis_h']:.1f} h · "
             f"{x['n_fit']} of {x['n_all']} sensors fitted · dots = data, lines = fits (grey = excluded sensors)",
             fontsize=9, color=INK2)
    fig.text(0.95, 0.955, f"page {k} / {n}", fontsize=8, color="#777", ha="right")
    ax = fig.add_axes([0.07, 0.38, 0.6, 0.52]); sensorgram(ax, x, inset=False)
    axz = fig.add_axes([0.73, 0.38, 0.24, 0.52]); sensorgram(axz, x, inset=False, legend=False, small=True)
    axz.set_xlim(-0.05, x["e"]["r"]["tb"] / 3600 + 1.0)
    axz.set_title("first part (zoom)", fontsize=9, color=INK2, loc="left")
    ka1_note = (f"{x['top']:g} nM slow phase: k = {x['k_slow']:.2e} 1/s ({100 * x['slow_share']:.0f}% of signal)"
                if x["ka1_ok"] else f"n.d.: at {x['top']:g} nM the slow component is the dimer binding")
    rows = [("Ka1  tetramer (1/Ms)", ka1_text(x), ka1_note),
            ("Kdis1 tetramer (1/s)", kdis1_text(x), f"half-life {t1_text(x)} days" +
             ("  (no decay detected in 15 h: bound)" if x["kdis1_bound"] else "")),
            ("KD1  tetramer (nM)", kd1_text(x), "Kdis1 / Ka1"),
            ("Ka2  dimer (1/Ms)", f"{x['ka2']:.3e}", "slope of k_obs vs [EphB2], all curves"),
            ("Kdis2 dimer (1/s)", f"{x['kdis2']:.3e}", f"half-life {LN2 / x['kdis2'] / 60:.1f} min"),
            ("KD2  dimer (nM)", f"{x['KD2_nM']:.3g}", "Kdis2 / Ka2"),
            ("R² association, top conc", f"{x['r2_top']:.4f}", "fit used for Ka1"),
            ("R² association, all curves", f"{x['r2_all']:.4f}", f"used for Ka2; k_obs regression R² = {x['r2_kobs']:.3f}"),
            ("R² dissociation", f"{x['r2_dis']:.4f}", "global fit, all curves (Kdis1, Kdis2)")]
    y = 0.29
    for name, val, note in rows:
        fig.text(0.07, y, name, fontsize=9, color=INK, family="DejaVu Sans Mono")
        fig.text(0.36, y, val, fontsize=9, color=INK, family="DejaVu Sans Mono", weight="bold")
        fig.text(0.5, y, note, fontsize=8, color=INK2)
        y -= 0.028
    pdf.savefig(fig); plt.close(fig)


def table_pages(pdf, runs):
    fig = plt.figure(figsize=(11, 8.5))
    fig.text(0.03, 0.955, "Extended runs (14-16 h dissociation): rate constants and fit quality", fontsize=13,
             weight="bold", color=INK)
    fig.text(0.03, 0.93, "1 = tetramer, 2 = dimer. Ka1 = (k_slow − Kdis1) / [EphB2]top from the top-concentration "
             "slow phase; Ka2 from all association curves; Kdis from the global dissociation fit.\n< / > = bound; "
             "n.d. = not determinable.", fontsize=7.8, color=INK2, va="top", linespacing=1.5)
    ax = fig.add_axes([0.01, 0.1, 0.98, 0.8]); ax.axis("off")
    head = ["GF", "Run", "Ligand", "Analyte", "Buffer", "Ka1\n(1/Ms)", "Kdis1\n(1/s)", "KD1\n(nM)", "Ka2\n(1/Ms)",
            "Kdis2\n(1/s)", "KD2\n(nM)", "R² assoc\ntop", "R² assoc\nall", "R²\ndissoc"]
    rows = [[str(x["gf"]), x["exp"].split()[0], x["lig"].replace("Ephrin", "E").replace("-Fc", ""), x["analyte"],
             x["buffer"], ka1_text(x), kdis1_text(x), kd1_text(x), f"{x['ka2']:.3e}", f"{x['kdis2']:.3e}",
             f"{x['KD2_nM']:.3g}", f"{x['r2_top']:.4f}", f"{x['r2_all']:.4f}", f"{x['r2_dis']:.4f}"] for x in runs]
    wd = [0.035, 0.065, 0.045, 0.085, 0.055, 0.085, 0.095, 0.06, 0.085, 0.085, 0.055, 0.065, 0.065, 0.06]
    tb = ax.table(cellText=rows, colLabels=head, loc="upper center", cellLoc="center", colWidths=[w / sum(wd) for w in wd])
    tb.auto_set_font_size(False); tb.set_fontsize(7.4); tb.scale(1, 1.62)
    for (i, j), c in tb.get_celld().items():
        c.set_edgecolor(GRID)
        if i == 0:
            c.set_facecolor("#f2f2ef"); c.set_text_props(weight="bold", color=INK)
        elif i in (9, 13):                                   # first row of each new ligand/buffer block
            c.visible_edges = "BRLT"
    fig.text(0.03, 0.065, "n.d. (2026.90, .96, .97): the top-concentration slow component matches the dimer k_obs expected "
             "from Ka2, so there is no separate second phase.\n< / > (2026.91, .97): no tetramer decay detected in 15 h, so Kdis1 "
             "is an upper bound. 2026.94: Ka2 regression R² 0.62, Ka2/KD2 approximate. 2026.90: 20-min association.",
             fontsize=7.4, color=INK2, linespacing=1.5)
    pdf.savefig(fig); plt.close(fig)

    fig = plt.figure(figsize=(11, 8.5))
    fig.text(0.04, 0.95, "EphB2-WT vs EphB2-R79P", fontsize=13, weight="bold", color=INK)
    fig.text(0.04, 0.935, "Student's t-test, two-tailed, equal variance. Only EphrinB3-Fc in PBST has replicates (n = 3 each);\n"
             "the PBSTB sets have one run per variant, shown below for reference (not tested). Buffers are not pooled.",
             fontsize=8.3, color=INK2, va="top", linespacing=1.5)
    ax2 = fig.add_axes([0.03, 0.5, 0.94, 0.38]); ax2.axis("off")
    grp = lambda g, lig="EphrinB3-Fc", buf="PBST": [x for x in runs if x["analyte"] == g and x["lig"] == lig and x["buffer"] == buf]
    wt, mu = grp("EphB2-WT"), grp("EphB2-R79P")
    rows2 = []
    for key, name, f in (("ka1", "Ka1 (1/Ms)", "{:.3e}"), ("kdis1", "Kdis1 (1/s)", "{:.3e}"), ("KD1_nM", "KD1 (nM)", "{:.3g}"),
                         ("ka2", "Ka2 (1/Ms)", "{:.3e}"), ("kdis2", "Kdis2 (1/s)", "{:.3e}"), ("KD2_nM", "KD2 (nM)", "{:.3g}")):
        a = [x[key] for x in wt]; b = [x[key] for x in mu]
        p = stats.ttest_ind(a, b, equal_var=True).pvalue
        rows2.append([name, f"{f.format(np.mean(a))} ± {f.format(np.std(a, ddof=1))}",
                      f"{f.format(np.mean(b))} ± {f.format(np.std(b, ddof=1))}", f"{np.mean(b) / np.mean(a):.2f}",
                      "<0.001" if p < 0.001 else f"{p:.3f}"])
    tb2 = ax2.table(cellText=rows2, colLabels=["EphrinB3-Fc, PBST", "WT mean ± SD (n=3)", "R79P mean ± SD (n=3)",
                                               "R79P / WT", "Student's p"], loc="upper center", cellLoc="center")
    tb2.auto_set_font_size(False); tb2.set_fontsize(8.4); tb2.scale(1, 1.8)
    ax3 = fig.add_axes([0.03, 0.08, 0.94, 0.36]); ax3.axis("off")
    rows3 = []
    for lig in ("EphrinB3-Fc", "EphrinB2-Fc"):
        a, b = grp("EphB2-WT", lig, "PBSTB")[0], grp("EphB2-R79P", lig, "PBSTB")[0]
        for key, name, f in (("kdis1", "Kdis1 (1/s)", None), ("KD1_nM", "KD1 (nM)", None), ("kdis2", "Kdis2 (1/s)", "{:.3e}"),
                             ("KD2_nM", "KD2 (nM)", "{:.3g}")):
            if key == "kdis1":
                va, vb = kdis1_text(a), kdis1_text(b)
            elif key == "KD1_nM":
                va, vb = kd1_text(a), kd1_text(b)
            else:
                va, vb = f.format(a[key]), f.format(b[key])
            rows3.append([f"{lig}, PBSTB", name, f"{va} ({a['exp'].split()[0]})", f"{vb} ({b['exp'].split()[0]})"])
    tb3 = ax3.table(cellText=rows3, colLabels=["Set", "Constant", "WT (single run)", "R79P (single run)"],
                    loc="upper center", cellLoc="center")
    tb3.auto_set_font_size(False); tb3.set_fontsize(8); tb3.scale(1, 1.45)
    for t in (tb2, tb3):
        for (i, j), c in t.get_celld().items():
            c.set_edgecolor(GRID)
            if i == 0:
                c.set_facecolor("#f2f2ef"); c.set_text_props(weight="bold", color=INK)
    pdf.savefig(fig); plt.close(fig)


def overview_png(runs):
    fig, axes = plt.subplots(4, 4, figsize=(21, 16))
    for ax, x in zip(axes.ravel(), runs):
        sensorgram(ax, x, inset=False, legend=False, small=True)
        ax.set_title(f"{x['lig'].replace('-Fc', '')} + {x['analyte']}  {x['exp'].split()[0]}  (GF {x['gf']}, {x['buffer']})",
                     fontsize=9, weight="bold", color=INK, loc="left")
    fig.suptitle("All 16 extended runs: full sensorgrams (association + 14-16 h dissociation), data = dots, fits = lines",
                 fontsize=13, weight="bold", color=INK, x=0.01, ha="left")
    fig.tight_layout(rect=[0, 0, 1, 0.97])
    fig.savefig(OUT_PNG, dpi=170); plt.close(fig)


# ------------------------------------------------------------------ workbook
def write_xlsx(runs):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    wb = Workbook()
    F, FB = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
    FILL = PatternFill("solid", fgColor="F2F2EF"); BOX = Border(bottom=Side(style="thin", color="D9D9D6"))

    def header(ws, cols, widths):
        for j, (h, w) in enumerate(zip(cols, widths), 1):
            c = ws.cell(row=1, column=j, value=h); c.font = FB; c.fill = FILL; c.border = BOX
            c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
            ws.column_dimensions[c.column_letter].width = w
        ws.row_dimensions[1].height = 58; ws.freeze_panes = "E2"

    ws = wb.active; ws.title = "Constants"
    cols = ["Global Fit #", "Run", "Ligand", "Analyte", "Buffer", "Association (min)", "Dissociation (h)",
            "Top concentration (nM)", "Slow-phase rate at top conc, k_slow (1/s)",
            "Fast-phase rate at top conc (1/s)", "Dimer k_obs expected at top conc = Ka2*[A]+Kdis2 (1/s)",
            "Ka1 tetramer (1/Ms)", "Kdis1 tetramer (1/s)", "Kdis1 type", "KD1 tetramer (nM)", "Ka2 dimer (1/Ms)",
            "Kdis2 dimer (1/s)", "KD2 dimer (nM)", "R² association, top concentration (Ka1)",
            "R² association, all curves (Ka2)", "R² k_obs vs [EphB2] regression (Ka2)",
            "R² dissociation, global fit (Kdis1, Kdis2)", "Tetramer half-life (days)", "Dimer half-life (min)", "Notes"]
    header(ws, cols, [9, 15, 12, 12, 8, 11, 11, 11, 14, 14, 16, 14, 14, 12, 11, 14, 14, 10, 13, 13, 13, 14, 12, 12, 70])
    fmts = {6: "0", 7: "0.0", 8: "0.0", 9: "0.000E+00", 10: "0.000E+00", 11: "0.000E+00", 12: "0.000E+00", 13: "0.000E+00",
            15: "0.000", 16: "0.000E+00", 17: "0.000E+00", 18: "0.00", 19: "0.0000", 20: "0.0000", 21: "0.000", 22: "0.0000",
            23: "0.00", 24: "0.0"}
    for i, x in enumerate(runs):
        rr = i + 2; x["row"] = rr
        v = {1: x["gf"], 2: x["exp"], 3: x["lig"], 4: x["analyte"], 5: x["buffer"], 6: round(x["assoc_min"], 1),
             7: round(x["dis_h"], 2), 8: x["top"], 9: float(f"{x['k_slow']:.10g}"), 10: float(f"{x['k_fast']:.10g}"),
             11: f"=P{rr}*H{rr}*1E-9+Q{rr}", 12: f"=(I{rr}-M{rr})/(H{rr}*1E-9)", 13: float(f"{x['kdis1']:.10g}"),
             14: "upper bound" if x["kdis1_bound"] else "measured", 15: f"=M{rr}/L{rr}*1E9",
             16: float(f"{x['ka2']:.10g}"), 17: float(f"{x['kdis2']:.10g}"), 18: f"=Q{rr}/P{rr}*1E9",
             19: float(f"{x['r2_top']:.10g}"), 20: float(f"{x['r2_all']:.10g}"), 21: float(f"{x['r2_kobs']:.10g}"),
             22: float(f"{x['r2_dis']:.10g}"), 23: f"=LN(2)/M{rr}/86400", 24: f"=LN(2)/Q{rr}/60", 25: notes(x)}
        if not x["ka1_ok"]:
            v[12] = v[15] = "n.d."
        for j, val in v.items():
            c = ws.cell(row=rr, column=j, value=val); c.font = F
            if j in fmts:
                c.number_format = fmts[j]
        ws.cell(row=rr, column=25).alignment = Alignment(wrap_text=True, vertical="top")

    st = wb.create_sheet("WT vs R79P")
    header(st, ["Constant (EphrinB3-Fc, PBST)", "WT mean", "WT SD", "WT n", "R79P mean", "R79P SD", "R79P n", "R79P / WT",
                "Student's t-test p (two-tailed, equal variance)"], [28, 13, 13, 7, 13, 13, 8, 10, 18])
    st.freeze_panes = "A2"
    wr = [x["row"] for x in runs if x["analyte"] == "EphB2-WT" and x["lig"] == "EphrinB3-Fc" and x["buffer"] == "PBST"]
    mr = [x["row"] for x in runs if x["analyte"] == "EphB2-R79P" and x["lig"] == "EphrinB3-Fc" and x["buffer"] == "PBST"]
    assert wr == list(range(wr[0], wr[-1] + 1)) and mr == list(range(mr[0], mr[-1] + 1)) and len(wr) == len(mr) == 3
    for i, (name, c, fmt) in enumerate((("Ka1 (1/Ms)", "L", "0.000E+00"), ("Kdis1 (1/s)", "M", "0.000E+00"),
                                        ("KD1 (nM)", "O", "0.000"), ("Ka2 (1/Ms)", "P", "0.000E+00"),
                                        ("Kdis2 (1/s)", "Q", "0.000E+00"), ("KD2 (nM)", "R", "0.00"))):
        rr = i + 2
        a, b = f"Constants!{c}{wr[0]}:{c}{wr[-1]}", f"Constants!{c}{mr[0]}:{c}{mr[-1]}"
        for j, val in enumerate([name, f"=AVERAGE({a})", f"=STDEV({a})", f"=COUNT({a})", f"=AVERAGE({b})",
                                 f"=STDEV({b})", f"=COUNT({b})", f"=E{rr}/B{rr}", f"=TTEST({a},{b},2,2)"], 1):
            st.cell(row=rr, column=j, value=val).font = F
        for col in (2, 3, 5, 6):
            st.cell(row=rr, column=col).number_format = fmt
        st.cell(row=rr, column=8).number_format = "0.00"; st.cell(row=rr, column=9).number_format = "0.000"
    st.cell(row=9, column=1, value=("WT = GF 69, 70, 73; R79P = GF 71, 72, 74. The PBSTB sets (EphrinB3: GF 67 vs 68; "
                                    "EphrinB2: GF 60 vs 61) have one run per variant and are not tested. Buffers are not pooled.")
            ).font = Font(name="Arial", size=9, italic=True)

    nt = wb.create_sheet("Notes"); nt.column_dimensions["A"].width = 130
    for i, t in enumerate([
        "All 16 extended runs (14-16 h dissociation) = the blue-highlighted rows of 'MASTER TABLE (BEST Colored)'. Generated 2026-09-30. "
        "1 = tetramer, 2 = dimer.",
        "Ka2: slope of k_obs vs [EphB2] from the per-sensor double-exponential association fits (all fitted curves).",
        "Kdis2, Kdis1: fast and slow rates of the global biexponential dissociation fit (shared rates, per-sensor amplitudes). "
        "Kdis1 type 'upper bound' = no tetramer decay detected in 15 h; KD1 and the tetramer half-life are then bounds (KD1 <, t1/2 >).",
        "Ka1: top-concentration association trace refit alone (1500 evenly spaced points), fast + slow exponential; "
        "Ka1 = (k_slow - Kdis1)/[EphB2]_top (column L), assuming the slow phase is a bimolecular step with rate linear in [EphB2] "
        "(only the top concentration shows it, so this is assumed, not tested).",
        "Ka1 is n.d. when the slow component, not the fast one, is closer to the dimer k_obs expected at that concentration "
        "(column K = Ka2*[A] + Kdis2): the 'slow' phase is then the dimer binding itself (2026.90, 2026.96, 2026.97).",
        "2026.90 (GF 59) has a 20-min association (16 h dissociation). It was previously skipped by a boundary rule that required "
        "the association to be >= 3% of the record; the rule is now >= 5 min of association, which changes no other file.",
        "KD = Kdis/Ka x 1e9 (nM). Half-lives = LN(2)/Kdis. R2 association (all curves) is pooled over all fitted sensors; "
        "R2 dissociation is the pooled global fit.",
        "Typed values are fit outputs; Ka1, KD1, KD2, the expected k_obs, the half-lives and the WT vs R79P statistics are formulas.",
        "Code: dimer_tetramer/extended_constants_all16.py"], 1):
        c = nt.cell(row=i, column=1, value=t); c.font = F; c.alignment = Alignment(wrap_text=True)
    wb.save(OUT_XLSX)


def main():
    runs = load()
    with PdfPages(OUT_PDF) as pdf:
        table_pages(pdf, runs)
        for k, x in enumerate(runs, 3):
            run_page(pdf, x, k, len(runs) + 2)
        pdf.infodict()["Title"] = "All 16 extended runs: rate constants and sensorgrams"
    overview_png(runs)
    write_xlsx(runs)
    cell, n = evaluate_workbook(OUT_XLSX)
    for x in runs:
        rr = x["row"]
        checks = [(f"R{rr}", x["KD2_nM"]), (f"K{rr}", x["k_exp"])]
        if x["ka1_ok"]:
            checks += [(f"L{rr}", x["ka1"]), (f"O{rr}", x["KD1_nM"])]
        for ref, want in checks:
            got = cell("Constants", ref)
            assert abs(got - want) <= 1e-7 * abs(want), (ref, got, want)
    wt = [x for x in runs if x["analyte"] == "EphB2-WT" and x["lig"] == "EphrinB3-Fc" and x["buffer"] == "PBST"]
    mu = [x for x in runs if x["analyte"] == "EphB2-R79P" and x["lig"] == "EphrinB3-Fc" and x["buffer"] == "PBST"]
    for i, key in enumerate(("ka1", "kdis1", "KD1_nM", "ka2", "kdis2", "KD2_nM")):
        p = stats.ttest_ind([x[key] for x in wt], [x[key] for x in mu], equal_var=True).pvalue
        assert abs(cell("WT vs R79P", f"I{i + 2}") - p) <= 1e-6 * p
    print(f"workbook: {n} formulas evaluated, all finite and matching Python")
    print(f"{'GF':3s} {'run':9s} {'lig':4s} {'analyte':11s} {'buf':6s} {'Ka1':>10s} {'Kdis1':>11s} {'KD1 nM':>8s} {'Ka2':>10s} "
          f"{'Kdis2':>10s} {'KD2 nM':>7s} {'R2top':>6s} {'R2all':>6s} {'R2dis':>6s}")
    for x in runs:
        print(f"{x['gf']:3d} {x['exp'].split()[0]:9s} {x['lig'][6:8]:4s} {x['analyte']:11s} {x['buffer']:6s} {ka1_text(x):>10s} "
              f"{kdis1_text(x):>11s} {kd1_text(x):>8s} {x['ka2']:10.3e} {x['kdis2']:10.3e} {x['KD2_nM']:7.3g} "
              f"{x['r2_top']:6.4f} {x['r2_all']:6.4f} {x['r2_dis']:6.4f}")
    for key in ("ka1", "kdis1", "KD1_nM", "ka2", "kdis2", "KD2_nM"):
        a = [x[key] for x in wt]; b = [x[key] for x in mu]
        print(f"  EB3 PBST {key:7s} WT {np.mean(a):.3g} ± {np.std(a, ddof=1):.2g}  R79P {np.mean(b):.3g} ± {np.std(b, ddof=1):.2g}  "
              f"p = {stats.ttest_ind(a, b, equal_var=True).pvalue:.3f}")
    print(OUT_PDF); print(OUT_PNG); print(OUT_XLSX)


if __name__ == "__main__":
    main()
